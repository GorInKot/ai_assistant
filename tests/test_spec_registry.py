"""Реестр позиций спецификаций по сканам (форма ГОСТ 21.110): правила и сквозной прогон."""

from __future__ import annotations

import base64
import io
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services import chat_documents as cd
from app.services import spec_registry as sr
from app.services.document_processing import DocumentError


def _line(name="", pos="", type_="", unit="", qty="", mass="", note="", conf=95.0, y=0.0):
    cells = [pos, name, type_, "", "", unit, qty, mass, note]
    return sr.TableLine(y=y, cells=cells, conf=[conf if c else 100.0 for c in cells])


# ------------------------------- единицы ----------------------------------


@pytest.mark.parametrize("raw, unit", [("шт.", "шт."), ("М", "м"), ("шт,", "шт."), ("шл.", "шт."), ("КМ", "км"), ("", "")])
def test_normalize_unit_known(raw, unit):
    assert sr.normalize_unit(raw) == (unit, True)


@pytest.mark.parametrize("raw", ["НЯ", "ил", "ви", "уу \\ми"])
def test_normalize_unit_unknown_kept_as_read(raw):
    value, sure = sr.normalize_unit(raw)
    assert not sure and value == raw.strip()


def test_single_letter_unit_needs_both_readings():
    # «т» и «л» — частые ошибки чтения «шт»: одного прочтения мало
    assert sr._agreed_unit(["ит", "т"]) == ""
    assert sr._agreed_unit(["т", "т"]) == "т"
    assert sr._agreed_unit(["м", "М"]) == "м"
    # многобуквенная, прочитанная точно, — достаточно одного прочтения
    assert sr._agreed_unit(["ил", "ШТ"]) == "шт"
    assert sr._agreed_unit(["НЯ", "НИ"]) == ""


# ---------------------------- сборка позиций ------------------------------


def test_groups_prefix_variants_and_sections():
    lines = [
        _line("Сети водоотведения наружные (С1)"),
        _line("Лоток железобетонный водоотводный Л 12.8", type_="ТУ 0000-002", unit="м", qty="64", mass="210,5"),
        _line("Труба полиэтиленовая напорная ПЭ100 SDR17 с защитной"),
        _line("оболочкой из полипропилена Д-160х9,5 Ду150", type_="ТУ 0000-001", unit="м", qty="87", mass="4,6*"),
        _line("D-225x13,4 Dy200", type_="ТУ 0000-001", unit="м", qty="143", mass="8,9*"),
        _line("Хомут ремонтный для труб ПЭ Д-160", unit="шт.", qty="18", mass="1,20"),
        _line("Д-225", unit="шт.", qty="31", mass="1,75"),
        _line("325х8,0 ГОСТ 10704-91 Труба стальная электросварная", unit="м", qty="6,5", mass="62,54"),
    ]
    items, section = sr.group_positions(lines, page=1)
    names = [it.cells[sr.COL_NAME] for it in items]
    assert names == [
        "Лоток железобетонный водоотводный Л 12.8",
        "Труба полиэтиленовая напорная ПЭ100 SDR17 с защитной оболочкой из полипропилена Д-160х9,5 Ду150",
        "Труба полиэтиленовая напорная ПЭ100 SDR17 с защитной оболочкой из полипропилена D-225х13,4 Dy200",
        "Хомут ремонтный для труб ПЭ Д-160",
        "Хомут ремонтный для труб ПЭ Д-225",
        # полное наименование, начинающееся с размера, — не вариант предыдущей позиции
        "325х8,0 ГОСТ 10704-91 Труба стальная электросварная",
    ]
    assert all(it.section == "Сети водоотведения наружные (С1)" for it in items)
    assert section == "Сети водоотведения наружные (С1)"
    assert [it.cells[sr.COL_QTY] for it in items] == ["64", "87", "143", "18", "31", "6,5"]


def test_quantity_on_first_line_then_continuation():
    lines = [
        _line("2. Материалы, поставляемые подрядчиком."),
        _line("2.1. Оборудование дозирования"),
        _line("Насос-дозатор плунжерный с регулятором хода", pos="2.1.1", type_="НД-25", unit="шт.", qty="6"),
        _line("двухголовочный, на давление 1,6 МПа,", type_="ТУ 0000-002"),
        _line("(с комплектом запасных частей)"),
        _line("Клапан предохранительный пружинный,", pos="2.1.2"),
        _line("с подрывным рычагом", unit="шт.", qty="3"),
    ]
    items, _ = sr.group_positions(lines, page=3)
    assert len(items) == 2
    first, second = items
    assert first.cells[sr.COL_POS] == "2.1.1"
    assert first.cells[sr.COL_NAME].endswith("(с комплектом запасных частей)")
    assert first.cells[sr.COL_TYPE] == "НД-25 ТУ 0000-002"
    assert first.section == "2. Материалы, поставляемые подрядчиком. / 2.1. Оборудование дозирования"
    # количество на второй строке позиции, начатой номером
    assert second.cells[sr.COL_NAME] == "Клапан предохранительный пружинный, с подрывным рычагом"
    assert second.cells[sr.COL_QTY] == "3"
    assert first.page == second.page == 3


def test_numbered_heading_is_not_a_continuation():
    lines = [
        _line("То же, но сечением 3х2,5 мм2", unit="км", qty="0,15"),
        _line("2.4. Крепёжные изделия"),
        _line("Скоба крепёжная однолапковая", unit="шт.", qty="40"),
    ]
    items, _ = sr.group_positions(lines, page=1)
    assert [it.cells[sr.COL_NAME] for it in items] == ["То же, но сечением 3х2,5 мм2", "Скоба крепёжная однолапковая"]
    assert items[1].section == "2.4. Крепёжные изделия"


def test_flags_instead_of_guessing():
    lines = [
        _line("Насос-дозатор", unit="шт.", qty="12 15"),  # зачёркнуто и вписано от руки
        _line("Переход 114х5", unit="НЯ", qty="3"),
        _line("ыы ree", conf=20.0),  # подпись или штамп
        _line("Провод", unit="км", qty="0,25", conf=40.0),
    ]
    lines[1].conf[sr.COL_UNIT] = 20.0
    lines[3].colored = True
    items, _ = sr.group_positions(lines, page=1)
    assert len(items) == 3
    assert "количество: не одно число (возможна правка)" in items[0].issues
    assert items[0].cells[sr.COL_QTY] == "12 15"  # не склеивается в «1215»
    assert "единица не из словаря" in items[1].issues
    assert "неуверенно прочитано: ед." in items[1].issues
    assert "рядом нераспознанный текст (возможна правка от руки)" in items[1].issues
    assert "цветные пометки у количества или массы (ручка, штамп)" in items[2].issues
    assert "неуверенно прочитано: ед., кол-во" in items[2].issues
    assert "наименование прочитано плохо" not in items[2].issues  # 40 — для наименования терпимо


def test_stamp_in_notes_does_not_flag_position():
    lines = [
        _line("Труба стальная 76х3,5", unit="м", qty="9"),
        _line(note="КОПИЯ", conf=20.0),  # штамп в «Примечании»
    ]
    items, _ = sr.group_positions(lines, page=1)
    assert len(items) == 1 and items[0].issues == []


def test_line_with_quantity_is_never_noise():
    line = _line("", unit="м", qty="7,5", mass="18,42", conf=10.0)
    assert not sr._is_noise(line)
    assert sr._is_noise(_line("^^^", conf=90.0))


# ------------------------------ сетка формы -------------------------------


def _form_lines(left: int, width: int, extra: list[int] = ()) -> list[sr.VLine]:
    xs = [int(left + share * width) for share in sr.FORM_SHARES] + list(extra)
    return [sr.VLine(x=x, y0=100, y1=2000, start=100) for x in sorted(xs)]


def test_match_form_by_proportions():
    lines = _form_lines(300, 4600, extra=[120, 180])  # слева линии боковой графы
    found = sr.match_form(lines, 4961)
    assert found is not None
    xs, matched, _used = found
    assert matched == 8 and xs[0] == 300 and xs[-1] == 4900


def test_match_form_rejects_other_tables():
    lines = [sr.VLine(x=300 + i * 460, y0=100, y1=2000, start=100) for i in range(11)]  # равные графы
    assert sr.match_form(lines, 4961) is None


def test_body_top_skips_title_row():
    mm = sr.MM
    # заголовок «СПЕЦИФИКАЦИЯ…» 14 мм, шапка 32 мм, строка номеров 8 мм, дальше строки по 8 мм
    rules = [int(y * mm) for y in (22, 54, 62, 70, 78, 86)]
    assert sr._body_top(rules, int(8 * mm)) == int(62 * mm)


# ------------------------------- Excel ------------------------------------


def _registry() -> sr.SpecRegistry:
    ok = sr.Position(page=1, section="Трубы")
    ok.cells = ["", "Труба стальная", "ГОСТ 10704-91", "", "", "м", "7,5", "18,42", ""]
    bad = sr.Position(page=2, section="Трубы", issues=["количество: не одно число (возможна правка)"])
    bad.cells = ["1.1", "=HYPERLINK(\"x\")", "", "", "", "шт.", "12 15", "", ""]
    doc = sr.SpecDocument(filename="a.pdf", code="123-ТМ-01-С-001", pages=3, table_pages=2,
                          skipped_pages=[3], positions=[ok, bad], seconds=4.2)
    return sr.SpecRegistry(documents=[doc], seconds=4.2)


def test_xlsx_numbers_flags_and_formula_safety():
    wb = load_workbook(io.BytesIO(sr.build_spec_xlsx(_registry())))
    assert wb.sheetnames == ["Позиции", "Документы", "Как читать"]
    ws = wb["Позиции"]
    assert [c.value for c in ws[1]] == sr.EXCEL_HEADERS
    assert ws.cell(2, 11).value == 7.5 and ws.cell(2, 12).value == 18.42  # числа, а не текст
    assert ws.cell(3, 11).value == "12 15"  # не число — остаётся как прочитано
    assert ws.cell(3, 16).value.startswith("количество")
    assert ws.cell(3, 1).fill.fgColor.rgb.endswith("FFF2B3")  # позиция к проверке — жёлтая
    assert ws.cell(2, 1).fill.fgColor.rgb in (None, "00000000")
    formula = ws.cell(3, 6)  # формула из скана не исполняется: текст с апострофом
    assert formula.data_type == "s" and formula.quotePrefix
    docs = wb["Документы"]
    assert [c.value for c in docs[2]][:7] == ["a.pdf", "123-ТМ-01-С-001", 3, 2, 2, 1, "3"]


def test_xlsx_marks_changes_subitems_and_excluded():
    doc = _registry().documents[0]
    sub = sr.Position(page=1, subitem=True)
    sub.cells = ["", "Блоки Б-600", "", "", "", "", "4.20", "", ""]
    gone = sr.Position(page=1, excluded=True, changes=["позиция зачёркнута (исключена изменением)"])
    gone.cells = ["", "Труба 76х4", "", "", "", "м", "12.6", "", ""]
    upd = sr.Position(page=1, changes=["кол-во: было 210.4"])
    upd.cells = ["", "Труба 133х4.5", "", "", "", "м", "198.7", "", ""]
    doc.positions = [sub, gone, upd]
    ws = load_workbook(io.BytesIO(sr.build_spec_xlsx(sr.SpecRegistry([doc], 1.0))))["Позиции"]
    assert ws.cell(2, 13).value == "да"
    assert ws.cell(3, 6).font.strike and ws.cell(3, 15).value.startswith("позиция зачёркнута")
    assert ws.cell(4, 15).value == "кол-во: было 210.4" and ws.cell(4, 1).fill.fgColor.rgb.endswith("DDEBF7")


# ------------------------------- чат --------------------------------------


@pytest.mark.parametrize("message", [
    "составь реестр спецификаций",
    "собери позиции из спецификации в Excel",
    "реестр позиций",
])
def test_chat_routes_to_spec_registry(monkeypatch, message):
    monkeypatch.setattr(sr, "unavailable_reason", lambda: None)
    monkeypatch.setattr(sr, "build_spec_registry", lambda files: _registry())
    result = cd.handle_chat_files(message, [("a.pdf", b"%PDF-1.4")], llm=None)
    assert result.attachment.filename == "specifications.xlsx"
    assert "Позиций: **2**" in result.answer and "Проверьте 1" in result.answer
    wb = load_workbook(io.BytesIO(base64.b64decode(result.attachment.content_base64)))
    assert wb["Позиции"].max_row == 3


def test_chat_spec_registry_explains_missing_ocr(monkeypatch):
    monkeypatch.setattr(sr, "unavailable_reason", lambda: "Распознавание сканов (Tesseract) на сервере не установлено")
    with pytest.raises(DocumentError, match="Tesseract"):
        cd.handle_chat_files("реестр спецификаций", [("a.pdf", b"%PDF-1.4")], llm=None)


# ----------------------- сквозной прогон по «скану» -----------------------

FONT = next((p for p in (
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
) if Path(p).exists()), None)

needs_ocr = pytest.mark.skipif(
    sr.unavailable_reason() is not None or FONT is None or not shutil.which("tesseract"),
    reason="нужны Tesseract, OpenCV и шрифт с кириллицей",
)

ROWS = [
    # позиция, наименование, тип, ед., кол-во, масса
    ("", "Трубы и детали трубопроводов", "", "", "", ""),
    ("1", "Труба стальная электросварная 114х5", "ГОСТ 10704-91", "м", "12,5", "13,44"),
    ("2", "Отвод крутоизогнутый 90-114х5", "ГОСТ 17375-2001", "шт.", "4", "1,9"),
    ("3", "Хомут ремонтный стальной Д-160", "ТУ 0000-001", "шт.", "18", "1,20"),
    ("", "Д-225", "ТУ 0000-001", "шт.", "31", "1,75"),
]


def _synthetic_scan(rotate: int = 0, content: list[tuple[str, ...]] = ROWS) -> bytes:
    """Лист A3 со спецификацией по форме ГОСТ 21.110, как картинка в PDF (скан без текста)."""
    import pymupdf
    from PIL import Image, ImageDraw, ImageFont

    mm = sr.MM
    img = Image.new("L", (int(420 * mm), int(297 * mm)), 255)
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT, int(3.5 * mm))
    small = ImageFont.truetype(FONT, int(2.5 * mm))
    left, width = 20 * mm, 395 * mm
    xs = [left + s * width for s in sr.FORM_SHARES]
    top, header, numbers = 10 * mm, 40 * mm, 48 * mm
    rows = [numbers + i * 8 * mm for i in range(1, 22)]
    bottom = rows[-1]
    for x in xs:
        draw.line([(x, top), (x, bottom)], fill=0, width=4)
    for y in [top, header, numbers, *rows]:
        draw.line([(xs[0], y), (xs[-1], y)], fill=0, width=4 if y in (top, header, numbers, bottom) else 2)
    titles = ["Поз.", "Наименование и техническая характеристика", "Тип, марка", "Код",
              "Завод", "Ед. изм.", "Кол.", "Масса", "Примечание"]
    for i, title in enumerate(titles):
        draw.text((xs[i] + 1.5 * mm, top + 12 * mm), title, font=small, fill=0)
        draw.text(((xs[i] + xs[i + 1]) / 2, header + 1.5 * mm), str(i + 1), font=small, fill=0)
    for r, (pos, name, type_, unit, qty, mass) in enumerate(content):
        y = numbers + r * 8 * mm + 2 * mm
        for col, text in ((0, pos), (1, name), (2, type_), (5, unit), (6, qty), (7, mass)):
            if text:
                draw.text((xs[col] + 2 * mm, y), text, font=font, fill=0)
    stamp_y = bottom + 3 * mm
    draw.rectangle([xs[-1] - 185 * mm, stamp_y, xs[-1], stamp_y + 15 * mm], outline=0, width=4)
    draw.text((xs[-1] - 150 * mm, stamp_y + 4 * mm), "12345678.001-Р-100.000.000-ТМ-01-С-001", font=font, fill=0)
    if rotate:
        img = img.rotate(rotate, expand=True, fillcolor=255)
    png = io.BytesIO()
    img.save(png, format="PNG")
    pdf = pymupdf.open()
    w_pt, h_pt = img.width / sr.DPI * 72, img.height / sr.DPI * 72
    page = pdf.new_page(width=w_pt, height=h_pt)
    page.insert_image(page.rect, stream=png.getvalue())
    return pdf.tobytes()


# Латинские и кириллические двойники (T/Т, C/С…) на скане неразличимы — сравниваем без них.
_LOOKALIKES = str.maketrans("TCHPMKXAEBOtcpkxaeo", "ТСНРМКХАЕВОтсркхаео")


@needs_ocr
@pytest.mark.parametrize("rotate", [0, 90])
def test_end_to_end_synthetic_scan(rotate):
    doc = sr.process_pdf(_synthetic_scan(rotate), "spec.pdf", page_budget=5)
    assert doc.error is None and doc.table_pages == 1
    got = [(p.cells[sr.COL_NAME].translate(_LOOKALIKES), p.cells[sr.COL_UNIT], p.cells[sr.COL_QTY], p.cells[sr.COL_MASS])
           for p in doc.positions]
    assert got == [
        (name.translate(_LOOKALIKES), unit, qty, mass) for name, unit, qty, mass in [
            ("Труба стальная электросварная 114х5", "м", "12,5", "13,44"),
            ("Отвод крутоизогнутый 90-114х5", "шт.", "4", "1,9"),
            ("Хомут ремонтный стальной Д-160", "шт.", "18", "1,20"),
            ("Хомут ремонтный стальной Д-225", "шт.", "31", "1,75"),
        ]
    ]
    assert all(p.section == "Трубы и детали трубопроводов" for p in doc.positions)
    assert "12345678" in doc.code


@needs_ocr
def test_page_without_form_is_reported():
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page(width=842, height=595)
    page.insert_text((72, 72), "ЛИСТ РЕГИСТРАЦИИ ИЗМЕНЕНИЙ", fontsize=20)
    doc = sr.process_pdf(pdf.tobytes(), "other.pdf", page_budget=5)
    assert doc.table_pages == 0 and doc.skipped_pages == [1] and not doc.positions


@pytest.mark.parametrize("raw, fixed", [
    ("Труба 114хб", "Труба 114х6"),
    ("Труба 325х8,0 ГОСТ", "Труба 325х8,0 ГОСТ"),
    ("Отвод 90-114хб ст.20", "Отвод 90-114х6 ст.20"),
    ("Полоса 7хЗ3", "Полоса 7х33"),
    ("Труба 114x5 и 76×3", "Труба 114х5 и 76х3"),  # латинская x и знак умножения → кириллическая х
    ("Фланец Ст20 плоский", "Фланец Ст20 плоский"),  # вне размеров буквы не трогаем
])
def test_fix_dimensions(raw, fixed):
    assert sr.fix_dimensions(raw) == fixed


@pytest.mark.parametrize("raw, code", [
    ("12345678.002-P-100.001.010-TK-01-C-001ГО", "12345678.002-Р-100.001.010-ТК-01-С-001"),
    ("12345678.001-Р-100.000.000-НВК-02-С-001", "12345678.001-Р-100.000.000-НВК-02-С-001"),
    ("12345678.002-P-100.001.020-HK-01-C-004AЕЕ", "12345678.002-Р-100.001.020-НК-01-С-004"),
])
def test_clean_code(raw, code):
    assert sr.clean_code(raw) == code


def test_code_regex_requires_number_start():
    text = "СПЕЦИФИКАЦИЯ12345678.001-Р-100.000.000-НВК-02-С-001 R-100_000_000-VK-02-S-001-RC01-F01.PDF"
    found = [m.group(0) for m in sr._CODE_RE.finditer(text)]
    assert found == ["12345678.001-Р-100.000.000-НВК-02-С-001"]


def test_parameters_under_position_are_not_sections():
    # насосы: параметры под позицией начинаются с заглавной, но это не заголовки
    lines = [
        _line("Насос центробежный консольный, с электродвигателем", pos="1", unit="Шт.", qty="2", mass="185"),
        _line("Подача - 25 м3/ч", type_="12345678.001-Р-100.000.000-"),
        _line("Напор - 32 м", type_="-ТХ-01-ОЛ-001"),
        _line("Взрывозащищённое исполнение (зона 2)"),
        _line("Насос дренажный погружной", pos="2", unit="Шт.", qty="1", mass="42"),
    ]
    items, section = sr.group_positions(lines, page=1)
    assert len(items) == 2 and section == ""
    assert items[0].cells[sr.COL_NAME].endswith("Напор - 32 м Взрывозащищённое исполнение (зона 2)")
    assert items[0].cells[sr.COL_TYPE] == "12345678.001-Р-100.000.000- -ТХ-01-ОЛ-001"


def test_header_like_line_becomes_section():
    heading = _line("Материалы вспомогательные")
    heading.header_like = True  # по центру графы
    lines = [
        _line("Шнур асбестовый 6мм", unit="м", qty="15"),
        heading,
        _line("Заглушка эллиптическая Ду 100", unit="шт", qty="2"),
    ]
    items, section = sr.group_positions(lines, page=1)
    assert [it.section for it in items] == ["", "Материалы вспомогательные"]
    assert items[0].cells[sr.COL_NAME] == "Шнур асбестовый 6мм"


def test_breakdown_lines_are_subitems():
    lines = [
        _line("Блоки газобетонные стеновые марки D500", pos="1", type_="ГОСТ 31360-2007", unit="м3", qty="12.40"),
        _line("на клеевом растворе"),
        _line("в том числе:"),
        _line("Б-600.300.200", qty="4.20"),
        _line("Б-600.250.200", qty="8.20"),
        _line("Сетка стальная плетёная", pos="3", unit="м2", qty="56.0"),
    ]
    items, _ = sr.group_positions(lines, page=1)
    assert [it.subitem for it in items] == [False, True, True, False]
    assert items[1].cells[sr.COL_NAME] == "Блоки газобетонные стеновые марки D500 на клеевом растворе Б-600.300.200"
    assert items[1].issues == []  # единица у разбивки — в позиции выше, это не ошибка
    assert items[3].cells[sr.COL_QTY] == "56.0" and not items[3].subitem


def test_struck_values_are_replaced_and_recorded():
    changed = _line("Труба 133х4.5 Сталь 20", type_="ГОСТ 10704-91", unit="м", qty="198.7", mass="14.26")
    changed.struck = {sr.COL_QTY: ["210.4"]}
    removed = _line("", type_="", unit="м", qty="", mass="")
    removed.struck = {sr.COL_NAME: ["Труба", "76х4", "Сталь", "20"], sr.COL_QTY: ["12.6"]}
    lost = _line("Труба 114х5", unit="м", qty="", mass="13.44")
    lost.struck = {sr.COL_QTY: ["35.0"]}
    items, _ = sr.group_positions([changed, removed, lost], page=1)
    assert items[0].cells[sr.COL_QTY] == "198.7" and items[0].changes == ["кол-во: было (по скану) 210.4"]
    assert items[0].issues == []
    assert items[1].excluded and items[1].cells[sr.COL_NAME] == "Труба 76х4 Сталь 20"
    assert items[1].cells[sr.COL_QTY] == "12.6"
    assert "количество зачёркнуто, нового значения не найдено" in items[2].issues


def test_is_struck_detects_line_through_middle():
    import numpy as np

    word = np.zeros((30, 120), dtype=bool)
    word[5:25, 10:20] = True  # «буквы»
    word[5:25, 60:70] = True
    assert not sr.is_struck(word)
    word[15, 2:118] = True  # штрих по середине
    assert sr.is_struck(word)
    underline = np.zeros((30, 120), dtype=bool)
    underline[28, 2:118] = True  # подчёркивание ниже середины — не зачёркивание
    assert not sr.is_struck(underline)


@needs_ocr
def test_api_spec_registry_without_llm(client):
    """Через чат: модель не нужна (в тестах её нет), реестр строится по скану."""
    from tests.conftest import register

    headers = register(client, "spec@test.local")
    resp = client.post(
        "/api/ask-files",
        data={"message": "составь реестр спецификаций"},
        files=[("files", ("spec.pdf", _synthetic_scan(), "application/pdf"))],
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["attachment"]["filename"] == "specifications.xlsx"
    assert "Позиций: **4**" in body["answer"]
    ws = load_workbook(io.BytesIO(base64.b64decode(body["attachment"]["content_base64"])))["Позиции"]
    assert [ws.cell(r, 11).value for r in range(2, 6)] == [12.5, 4, 18, 31]


def test_fully_struck_line_does_not_crash():
    line = _line()  # все графы пусты: каждое слово зачёркнуто
    line.struck = {sr.COL_NAME: ["Труба", "76х4"], sr.COL_QTY: ["12.6"]}
    items, _ = sr.group_positions([_line("Труба 48х3", unit="м", qty="6.0"), line], page=1)
    assert len(items) == 2 and items[1].excluded


@pytest.mark.parametrize("raw, unit", [("м/м2", "м/м2"), ("шт./м2", "шт./м2"), ("мЗ", "м3"), ("пара", "пара")])
def test_normalize_compound_units(raw, unit):
    assert sr.normalize_unit(raw) == (unit, True)


def test_extra_columns_mean_other_form():
    # ведомость оборудования: 18 граф — среди них находятся «подходящие», но лишних слишком много
    import numpy as np

    h, w = 2000, 4961
    lines = _form_lines(300, 4600) + [sr.VLine(x=x, y0=100, y1=1900, start=100) for x in (1100, 1300, 1500, 2300)]
    vert = np.zeros((h, w), dtype=np.uint8)
    for ln in lines:
        vert[ln.y0:ln.y1, ln.x - 2:ln.x + 3] = 255
    gray = np.full((h, w), 255, dtype=np.uint8)
    gray[vert > 0] = 0
    with pytest.raises(sr.SpecFormNotFound, match="не форма спецификации"):
        sr.find_grid(gray)


def test_color_cast_is_not_pen_marks():
    pymupdf = pytest.importorskip("pymupdf")
    pdf = pymupdf.open()
    page = pdf.new_page(width=400, height=300)
    page.draw_rect(pymupdf.Rect(0, 0, 400, 300), color=None, fill=(1, 1, 1))
    for i in range(20):  # «текст» с жёлтым оттенком по всему листу
        page.draw_line((20, 20 + i * 12), (380, 20 + i * 12), color=(0.45, 0.35, 0.05), width=2)
    _gray, colored = sr.render_page(page)
    assert not colored.any()
    page.draw_rect(pymupdf.Rect(200, 100, 230, 120), color=(0, 0, 1), fill=(0, 0, 1))  # синяя ручка
    page2 = pdf.new_page(width=400, height=300)
    for i in range(20):
        page2.draw_line((20, 20 + i * 12), (380, 20 + i * 12), color=(0, 0, 0), width=2)
    page2.draw_rect(pymupdf.Rect(200, 100, 230, 120), color=(0.1, 0.2, 0.6), fill=(0.1, 0.2, 0.6))  # синяя ручка
    _gray, colored = sr.render_page(page2)
    assert colored.any()


def _glyphs(pattern: str) -> "np.ndarray":
    import numpy as np

    rows = [list(r) for r in pattern.strip().splitlines()]
    return np.array([[ch == "#" for ch in r] for r in rows]).repeat(4, axis=0).repeat(4, axis=1)


SHT = """
#.#.#.###.
#.#.#..#..
#.#.#..#..
#####..#.#
"""
M = """
#...#
##.##
#.#.#
#...#
"""


def test_unit_resolved_by_template_from_same_document():
    sure = sr.Fragment(sr.COL_UNIT, 10, 16, "шт.", sr.AGREED_CONFIDENCE, image=_glyphs(SHT))
    meter = sr.Fragment(sr.COL_UNIT, 20, 16, "м", sr.AGREED_CONFIDENCE, image=_glyphs(M))
    unsure = sr.Fragment(sr.COL_UNIT, 30, 16, "ил.", 20.0, image=_glyphs(SHT))  # та же надпись, прочитана плохо
    odd = sr.Fragment(sr.COL_UNIT, 40, 16, "НЯ", 20.0, image=_glyphs("##\n##\n.#\n#."))  # ни на что не похожа
    templates: dict = {}
    sr.resolve_units_by_templates([sure, meter, unsure, odd], templates)
    assert unsure.text == "шт." and unsure.conf == sr.AGREED_CONFIDENCE
    assert odd.text == "НЯ" and odd.conf == 20.0  # не угадываем
    assert set(templates) == {"шт", "м"}


def test_rare_unit_one_letter_off_is_flagged():
    items = []
    for i in range(12):
        p = sr.Position(page=1)
        p.cells = ["", f"Труба {i}", "", "", "", "м", "1", "", ""]
        items.append(p)
    odd = sr.Position(page=1)
    odd.cells = ["", "Труба 133х4.5", "", "", "", "км", "198.7", "", ""]
    other = sr.Position(page=1)
    other.cells = ["", "Кольцо", "", "", "", "шт.", "4", "", ""]
    sr.flag_unit_outliers(items + [odd, other])
    assert odd.issues == ["единица «км» редкая для документа (обычно «м»)"]
    assert other.issues == []  # другая единица, а не «почти основная» — не ошибка чтения
