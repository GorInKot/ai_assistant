"""Реестр документов: PDF/сканы → Excel с заданными полями и проверкой по тексту."""

from __future__ import annotations

import base64
import io

import pymupdf
import pytest
from openpyxl import load_workbook

from app.services import chat_documents as cd
from app.services import document_processing as dp
from app.services import registry as rg
from tests.conftest import register


def _act_pdf(number: str = "15", company: str = "ООО «Вектор»", total: str = "1 250 000,00") -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_htmlbox(
        pymupdf.Rect(60, 50, 540, 300),
        f"<h3>Акт выполненных работ № {number}</h3>"
        f"<p>от 15 ноября 2026 г.</p><p>Подрядчик: {company}</p>"
        f"<p>Итого к оплате: {total} руб.</p>",
    )
    return doc.tobytes()


class _FakeLLM:
    """Возвращает заранее заданные значения по порядку документов."""

    def __init__(self, answers: list[dict]) -> None:
        self.answers = list(answers)
        self.calls: list[tuple[str, list[dict]]] = []

    def extract_fields(self, text, fields):
        self.calls.append((text, fields))
        data = self.answers.pop(0)
        return {f["key"]: data.get(f["key"]) for f in fields}


# ----------------------- поля из сообщения -----------------------------


@pytest.mark.parametrize(
    "message, expected",
    [
        ("составь реестр: номер, дата, подрядчик, сумма", ["Номер", "Дата", "Подрядчик", "Сумма"]),
        ("Реестр актов. Поля: номер акта; дата; сумма", ["Номер акта", "Дата", "Сумма"]),
        ("извлеки номер, дату и сумму из каждого файла", ["Номер", "Дату", "Сумму"]),
        ("реестр: номер, дата, в excel", ["Номер", "Дата"]),
        ("реестр: Номер, номер, НОМЕР", ["Номер"]),
    ],
)
def test_parse_fields(message, expected):
    assert rg.parse_fields_from_message(message)[: len(expected)] == expected


@pytest.mark.parametrize("message", ["", "составь реестр", "в excel", "что тут"])
def test_parse_fields_none_when_not_specified(message):
    assert rg.parse_fields_from_message(message) is None


def test_parse_fields_limits():
    many = ", ".join(f"поле{i}" for i in range(30))
    assert len(rg.parse_fields_from_message("реестр: " + many)) == rg.MAX_REGISTRY_FIELDS


# --------------------- проверка значений по тексту ---------------------


TEXT = "Акт № 15 от 15 ноября 2026 г. Подрядчик: ООО «Вектор». Итого: 1 250 000,00 руб."


@pytest.mark.parametrize(
    "value",
    ["15", "ООО Вектор", "Вектор", "1250000", "1 250 000,00", "15.11.2026", "15.11.26", "Акт"],
)
def test_grounded_values(value):
    assert rg.is_grounded(value, TEXT)


@pytest.mark.parametrize(
    "value",
    ["99", "ООО Ромашка", "2 000 000", "16.11.2026", "01.01.2020", "Договор поставки"],
)
def test_hallucinated_values_not_grounded(value):
    assert not rg.is_grounded(value, TEXT)


def test_free_text_fields_are_not_verified():
    assert rg.is_free_text_field("Краткое содержание")
    assert rg.is_free_text_field("Предмет договора")
    assert not rg.is_free_text_field("Номер")


# --------------------------- реестр целиком ---------------------------


def test_registry_flags_values_missing_from_document():
    data = dp.parse_pdf(_act_pdf(), "act.pdf")
    llm = _FakeLLM([{"Номер": "15", "Подрядчик": "ООО «Ромашка»", "Сумма": "1 250 000,00"}])
    result = rg.build_registry([data], ["Номер", "Подрядчик", "Сумма", "Краткое содержание"], llm)
    row = result.rows[0]
    assert row.values["Номер"] == "15"
    assert row.unverified == ["Подрядчик"]  # выдумка модели пойманa
    assert result.unverified_cells == 1


def test_registry_row_without_text_is_reported_not_crashed(monkeypatch):
    monkeypatch.setattr(dp.shutil, "which", lambda name: None)
    from tests.test_pdf_conversion import _scan_pdf

    data = dp.parse_pdf(_scan_pdf("Invoice 12345"), "scan.pdf")
    llm = _FakeLLM([])
    result = rg.build_registry([data], ["Номер"], llm)
    assert "OCR" in result.rows[0].error and llm.calls == []


def test_registry_limit_on_documents():
    docs = [dp.parse_pdf(_act_pdf(), f"{i}.pdf") for i in range(rg.MAX_REGISTRY_DOCS + 1)]
    with pytest.raises(dp.DocumentError, match="Лимит"):
        rg.build_registry(docs, ["Номер"], _FakeLLM([]))


def test_lookup_tolerates_case_and_spacing():
    assert rg._lookup({"номер  акта": "7"}, "Номер акта") == "7"


def test_xlsx_highlights_unverified_and_reports_errors():
    row_ok = rg.RegistryRow("a.pdf", {"Номер": "15", "Подрядчик": "ООО Ромашка"}, unverified=["Подрядчик"], pages=1)
    row_bad = rg.RegistryRow("b.pdf", {"Номер": "", "Подрядчик": ""}, error="в документе не найден текст")
    xlsx = rg.build_registry_xlsx(rg.RegistryResult(["Номер", "Подрядчик"], [row_ok, row_bad], 1.0, False))
    ws = load_workbook(io.BytesIO(xlsx)).active
    assert [c.value for c in ws[1]][:3] == ["Файл", "Номер", "Подрядчик"]
    assert ws["C2"].fill.fgColor.rgb.endswith("FFF2B3")  # выдумка подсвечена
    assert ws["B2"].fill.fgColor.rgb in ("00000000", "0")  # проверенное — нет
    assert ws["D2"].value == "Подрядчик"  # колонка «Проверить»
    assert "НЕ РАЗОБРАН" in ws["D3"].value


# ------------------------- маршрутизация в чате -------------------------


def _files(n: int):
    return [(f"act{i}.pdf", _act_pdf(number=str(10 + i))) for i in range(n)]


def test_several_pdfs_without_instruction_make_registry_with_default_fields():
    llm = _FakeLLM([{"Номер": "10"}, {"Номер": "11"}])
    res = cd.handle_chat_files("", _files(2), llm)
    assert res.attachment.filename == "registry.xlsx"
    assert "Поля вы не указали" in res.answer
    ws = load_workbook(io.BytesIO(base64.b64decode(res.attachment.content_base64))).active
    assert [c.value for c in ws[1]][1:7] == rg.DEFAULT_FIELDS
    assert len(llm.calls) == 2


def test_registry_uses_fields_from_message():
    llm = _FakeLLM([{"Номер": "10", "Сумма": "1 250 000,00"}])
    res = cd.handle_chat_files("составь реестр: номер, сумма", _files(1), llm)
    assert "Поля: Номер, Сумма." in res.answer
    assert [f["key"] for f in llm.calls[0][1]] == ["Номер", "Сумма"]
    assert "Поля вы не указали" not in res.answer


def test_single_pdf_without_keywords_still_converts():
    class NoLLM:
        def extract_fields(self, *a, **k):
            raise AssertionError("LLM не нужен для конвертации")

    res = cd.handle_chat_files("", _files(1), NoLLM())
    assert res.attachment.filename == "act0.docx"


def test_registry_warns_about_unverified_values_in_chat():
    llm = _FakeLLM([{"Номер": "777"}])
    res = cd.handle_chat_files("реестр: номер", _files(1), llm)
    assert "Проверьте" in res.answer and "жёлтым" in res.answer


def test_registry_too_many_documents_rejected_before_any_llm_call():
    llm = _FakeLLM([])
    with pytest.raises(dp.DocumentError, match="Лимит"):
        cd.handle_chat_files("реестр: номер", _files(rg.MAX_REGISTRY_DOCS + 1), llm)
    assert llm.calls == []


def test_endpoint_builds_registry(client, monkeypatch):
    import app.state as state

    monkeypatch.setattr(
        state.llm_service, "extract_fields", lambda text, fields: {f["key"]: "15" for f in fields}
    )
    headers = register(client, "reg@test.local")
    resp = client.post(
        "/api/ask-files",
        data={"message": "реестр: номер"},
        files=[("files", (f"a{i}.pdf", _act_pdf(), "application/pdf")) for i in range(2)],
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["attachment"]["filename"] == "registry.xlsx"


def test_endpoint_returns_503_when_llm_unavailable(client):
    headers = register(client, "reg2@test.local")  # в тестах LLM не настроен
    resp = client.post(
        "/api/ask-files",
        data={"message": "реестр: номер"},
        files=[("files", ("a.pdf", _act_pdf(), "application/pdf"))],
        headers=headers,
    )
    assert resp.status_code == 503


# ------------- исправления по результатам прогона на реальной модели -------------


@pytest.mark.parametrize(
    "value, text, expected",
    [
        ("412800.0", "Итого к оплате 412 800,00 руб.", "412 800,00"),
        ("3 480 000.00 руб.", "Стоимость работ составляет 3 480 000,00 руб., в том числе НДС", "3 480 000,00"),
        ("1250000", "сумма 1 250 000 руб", "1 250 000"),
        ("99", "номер 15 и сумма 100", "99"),  # в тексте нет — не трогаем
        ("А-218", "счёт А-218 от 03.10.2026", "А-218"),  # не число
        ("15", "Акт № 15 от 15 ноября", "15"),
    ],
)
def test_snap_number_to_source(value, text, expected):
    assert rg.snap_number_to_source(value, text) == expected


def test_flatten_dict_and_list_values_from_model():
    from app.llm import LLMService

    flat = LLMService._flatten_value
    assert flat({"Заказчик": "ООО А", "Подрядчик": "ООО Б"}) == "Заказчик: ООО А; Подрядчик: ООО Б"
    assert flat(["ООО А", "ООО Б"]) == "ООО А; ООО Б"
    assert flat(None) == "" and flat(15) == "15" and flat("  x ") == "x"
    assert "{" not in flat({"a": {"b": "c"}})


def test_registry_sends_field_hints_for_default_fields_only():
    llm = _FakeLLM([{}, {}])
    rg.build_registry([dp.parse_pdf(_act_pdf(), "a.pdf")], None, llm)
    spec = {f["key"]: f["description"] for f in llm.calls[0][1]}
    assert "акт, счёт, письмо" in spec["Тип документа"]

    rg.build_registry([dp.parse_pdf(_act_pdf(), "b.pdf")], ["Прораб"], llm)
    assert llm.calls[1][1] == [{"key": "Прораб", "title": "Прораб", "description": None}]


def test_registry_value_is_snapped_to_source_notation():
    data = dp.parse_pdf(_act_pdf(total="412 800,00"), "act.pdf")
    llm = _FakeLLM([{"Сумма": "412800.0"}])
    row = rg.build_registry([data], ["Сумма"], llm).rows[0]
    assert row.values["Сумма"] == "412 800,00" and row.unverified == []


def test_number_prefix_is_stripped():
    data = dp.parse_pdf(_act_pdf(number="47"), "act.pdf")
    row = rg.build_registry([data], ["Номер"], _FakeLLM([{"Номер": "№ 47"}])).rows[0]
    assert row.values["Номер"] == "47" and row.unverified == []
