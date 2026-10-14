"""Оценка ответов и аналитика пробелов в знаниях: журнал вопросов, отчёт, выгрузка, права."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

import app.state as state
from app.db import QaLog, SessionLocal
from app.kb import KnowledgeBaseIndex
from app.services import ask_service, gaps
from tests.conftest import ADMIN_EMAIL, ADMIN_PASSWORD, auth_headers, register

KB_TEXT = """# Статусы заявки в ЕКТП

Черновик — статус заявки после создания, до размещения на обработку.
На обработке — заявка размещена и ожидает проверки диспетчером.
"""

ANSWER = "Заявка в ЕКТП после создания получает статус Черновик."


@pytest.fixture
def kb(tmp_path: Path, monkeypatch):
    folder = tmp_path / "ЕКТП_Транспорт"
    folder.mkdir()
    (folder / "Статусы.md").write_text(KB_TEXT, encoding="utf-8")
    index = KnowledgeBaseIndex(tmp_path)
    index.build()
    monkeypatch.setattr(ask_service, "kb_index", index)
    monkeypatch.setattr(state.llm_service, "rerank_results", lambda q, cands, top_n=16: list(cands)[:top_n])
    monkeypatch.setattr(state.llm_service, "generate_answer", lambda q, ctx, intent="procedure": ANSWER)
    return index


def _ask(client, headers, question, conversation_id):
    resp = client.post("/api/ask", headers=headers, json={"question": question, "conversation_id": conversation_id})
    assert resp.status_code == 200, resp.text
    return resp.json()


def _conversation(client, headers):
    return client.post("/api/conversations", headers=headers, json={}).json()["id"]


def _admin(client):
    return auth_headers(client, ADMIN_EMAIL, ADMIN_PASSWORD)


def _seed(db, **kwargs):
    defaults = dict(question="вопрос", kind="missing", user_id=1, division="ЦА", sources_json="[]")
    db.add(QaLog(**{**defaults, **kwargs}))


# ------------------------------ оценка ответа ------------------------------


def test_answer_response_carries_message_id_and_rating_roundtrip(client, kb):
    headers = register(client, "r1@test.local")
    conv = _conversation(client, headers)
    body = _ask(client, headers, "Какие статусы бывают у заявки в ЕКТП?", conv)
    assert body["kind"] == "answer" and isinstance(body["message_id"], int)
    assert "resolved_question" not in body  # служебное поле не уходит наружу

    url = f"/api/messages/{body['message_id']}/rating"
    assert client.put(url, headers=headers, json={"rating": 1}).json()["rating"] == 1
    history = client.get(f"/api/conversations/{conv}", headers=headers).json()
    assert [m["rating"] for m in history["messages"] if m["role"] == "assistant"] == [1]

    client.put(url, headers=headers, json={"rating": -1})
    client.put(url, headers=headers, json={"rating": None})  # снять оценку
    history = client.get(f"/api/conversations/{conv}", headers=headers).json()
    assert [m["rating"] for m in history["messages"] if m["role"] == "assistant"] == [None]


def test_rating_rejects_invalid_values(client, kb):
    headers = register(client, "r2@test.local")
    body = _ask(client, headers, "Какие статусы бывают у заявки в ЕКТП?", _conversation(client, headers))
    resp = client.put(f"/api/messages/{body['message_id']}/rating", headers=headers, json={"rating": 5})
    assert resp.status_code == 422


def test_cannot_rate_someone_elses_message(client, kb):
    owner = register(client, "owner@test.local")
    other = register(client, "other@test.local")
    body = _ask(client, owner, "Какие статусы бывают у заявки в ЕКТП?", _conversation(client, owner))
    resp = client.put(f"/api/messages/{body['message_id']}/rating", headers=other, json={"rating": -1})
    assert resp.status_code == 404


def test_cannot_rate_user_message(client, kb):
    headers = register(client, "r3@test.local")
    conv = _conversation(client, headers)
    _ask(client, headers, "Какие статусы бывают у заявки в ЕКТП?", conv)
    detail = client.get(f"/api/conversations/{conv}", headers=headers).json()
    user_msg = next(m for m in detail["messages"] if m["role"] == "user")
    resp = client.put(f"/api/messages/{user_msg['id']}/rating", headers=headers, json={"rating": 1})
    assert resp.status_code == 404


def test_rating_requires_auth(client):
    assert client.put("/api/messages/1/rating", json={"rating": 1}).status_code == 401


# ------------------------------ журнал вопросов ------------------------------


def test_questions_are_logged_with_kind_and_rating_is_mirrored(client, kb):
    headers = register(client, "l1@test.local", division="ЦА")
    conv = _conversation(client, headers)
    answered = _ask(client, headers, "Какие статусы бывают у заявки в ЕКТП?", conv)
    refused = _ask(client, headers, "Какая погода на Марсе?", conv)
    _ask(client, headers, "Что ты умеешь?", conv)  # smalltalk в журнал не попадает
    client.put(f"/api/messages/{refused['message_id']}/rating", headers=headers, json={"rating": -1})

    with SessionLocal() as db:
        rows = db.query(QaLog).order_by(QaLog.id).all()
    assert [r.kind for r in rows] == ["answer", "missing"]
    assert rows[0].message_id == answered["message_id"] and rows[0].division == "ЦА"
    assert rows[0].claims_total == 1 and rows[0].claims_unsupported == 0
    assert "Статусы.md" in rows[0].sources_json
    assert rows[1].rating == -1 and rows[0].rating is None


def test_smalltalk_is_not_logged_as_a_question(client, kb):
    headers = register(client, "l2@test.local")
    resp = client.post("/api/ask", headers=headers, json={"question": "Привет"})
    assert resp.json()["kind"] == "smalltalk"
    with SessionLocal() as db:
        assert db.query(QaLog).count() == 0


def test_log_survives_conversation_deletion(client, kb):
    headers = register(client, "l3@test.local")
    conv = _conversation(client, headers)
    _ask(client, headers, "Какая погода на Марсе?", conv)
    client.delete(f"/api/conversations/{conv}", headers=headers)
    with SessionLocal() as db:
        assert db.query(QaLog).count() == 1


# ------------------------------ объединение вопросов ------------------------------


def test_lexical_clustering_merges_paraphrases_and_separates_topics():
    qs = [
        "Как оформить командировку за границу?",
        "Оформление командировки",
        "Как оформить отпуск?",
        "Командировка за границу: порядок оформления",
    ]
    clusters, used_embeddings = gaps.cluster_questions(qs, embedder=None)
    assert used_embeddings is False
    groups = sorted(sorted(c) for c in clusters)
    assert groups == [[0, 1, 3], [2]]


class FakeEmbedder:
    """Вектор по теме: слова «виза»/«загранкомандировка» — одно направление."""

    def embed_documents(self, texts):
        out = []
        for t in texts:
            low = t.lower()
            if "виз" in low or "загран" in low:
                out.append([1.0, 0.1, 0.0])
            else:
                out.append([0.0, 0.1, 1.0])
        return np.array(out, dtype=np.float32)


def test_embedding_clustering_catches_paraphrases_without_shared_words():
    qs = ["Нужна ли виза сотруднику?", "Загранкомандировка: порядок", "Как заказать пропуск?"]
    clusters, used_embeddings = gaps.cluster_questions(qs, embedder=FakeEmbedder())
    assert used_embeddings is True
    assert sorted(sorted(c) for c in clusters) == [[0, 1], [2]]


def test_same_template_different_subject_is_not_merged():
    # «Как оформить …» / «Как заказать … на объект» — общий шаблон, разные потребности
    qs = [
        "Как оформить командировку?",
        "Как оформить отпуск?",
        "Как заказать пропуск на объект?",
        "Как заказать транспорт на объект?",
    ]
    clusters, _ = gaps.cluster_questions(qs, embedder=None)
    assert len(clusters) == 4


def test_close_embeddings_alone_do_not_merge_template_questions():
    class Close:
        """Близость 0.77 (как у настоящей модели для шаблонных вопросов), общих слов нет."""

        def embed_documents(self, texts):
            a = np.array([1.0, 0.0], dtype=np.float32)
            b = np.array([0.77, 0.638], dtype=np.float32)  # cos(a, b) = 0.77
            return np.array([a if "командировку" in t else b for t in texts])

    clusters, used = gaps.cluster_questions(["Как оформить командировку?", "Как оформить отпуск?"], Close())
    assert used is True and len(clusters) == 2


def test_broken_embedder_falls_back_to_words():
    class Broken:
        def embed_documents(self, texts):
            raise RuntimeError("Ollama недоступна")

    clusters, used_embeddings = gaps.cluster_questions(["Оформление командировки", "Как оформить командировку"], Broken())
    assert used_embeddings is False and len(clusters) == 1


# ------------------------------ отчёт ------------------------------


def test_report_separates_no_answer_from_weak_document_and_counts_people():
    with SessionLocal() as db:
        # Тема без ответа: один сотрудник переформулировал трижды + второй спросил то же
        _seed(db, question="Как оформить командировку за границу?", user_id=1, division="ЦА")
        _seed(db, question="Оформление командировки за границу", user_id=1, division="ЦА")
        _seed(db, question="Порядок оформления командировки за границу", user_id=1, division="ЦА")
        _seed(db, question="Командировка за границу как оформить", user_id=2, division="Филиал")
        # Документ есть, но не помогает
        _seed(db, question="Сколько дней отпуск", kind="answer", rating=-1, user_id=3,
              sources_json='[{"relative_path": "ЦУС/Отпуск.md", "source_type": "context"}]')
        # Не проблема: ответ, который оценили «полезно»
        _seed(db, question="Какие статусы у заявки", kind="answer", rating=1, user_id=4)
        db.commit()

        report = gaps.build_report(db, days=30)

    s = report["summary"]
    assert s["questions"] == 6 and s["answered"] == 2 and s["refused"] == 4 and s["clarifications"] == 0
    assert s["rated"] == 2 and s["useful"] == 1 and s["usefulness"] == 0.5

    by_problem = {g["problem"]: g for g in report["gaps"]}
    missing = by_problem["no_document"]
    assert missing["requests"] == 4 and missing["people"] == 2  # людей, а не обращений
    assert missing["divisions"] == ["Филиал", "ЦА"]
    weak = by_problem["weak_document"]
    assert weak["not_useful"] == 1 and weak["sources"][0]["relative_path"] == "ЦУС/Отпуск.md"
    assert len(report["gaps"]) == 2  # удачный ответ пробелом не считается
    assert report["grouping"] == "words"


def test_refusal_rated_useful_is_still_a_gap():
    # «Полезно» на «в базе нет информации» — благодарность за честность, а не отсутствие пробела
    with SessionLocal() as db:
        _seed(db, question="Что такое ЕКТС?", kind="missing", rating=1)
        db.commit()
        report = gaps.build_report(db, days=30)
    assert report["summary"]["questions"] == 1 and report["summary"]["refused"] == 1
    assert [g["topic"] for g in report["gaps"]] == ["Что такое ЕКТС?"]


def test_clarification_dead_end_and_rated_down_are_gaps_but_resolved_one_is_not():
    with SessionLocal() as db:
        _seed(db, question="что такое GigaChat?", kind="clarify", user_id=1, rating=-1)
        _seed(db, question="кто согласует акт приёмки", kind="clarify", user_id=2)  # ушёл без ответа
        _seed(db, question="как получить допуск на объект", kind="clarify", user_id=3)
        _seed(db, question="как получить допуск на объект подрядчику", kind="answer", user_id=3)  # уточнил
        db.commit()
        report = gaps.build_report(db, days=30)

    topics = sorted(g["topic"] for g in report["gaps"])
    assert topics == ["кто согласует акт приёмки", "что такое GigaChat?"]
    s = report["summary"]
    assert s["questions"] == 4 and s["clarifications"] == 3
    assert all(g["problem"] == "no_document" for g in report["gaps"])  # нет ответа, а не плохой документ


def test_clarification_followed_up_too_late_is_still_a_dead_end():
    from datetime import datetime, timedelta, timezone

    now = datetime.now(timezone.utc)
    with SessionLocal() as db:
        _seed(db, question="как получить допуск на объект", kind="clarify", user_id=3,
              created_at=now - timedelta(hours=2))
        _seed(db, question="совсем другой вопрос", kind="answer", user_id=3, created_at=now)
        db.commit()
        report = gaps.build_report(db, days=30)
    assert [g["topic"] for g in report["gaps"]] == ["как получить допуск на объект"]


def test_report_has_no_usefulness_without_ratings():
    with SessionLocal() as db:
        _seed(db, question="Погода на Марсе")
        db.commit()
        assert gaps.build_report(db, days=30)["summary"]["usefulness"] is None


def test_report_ignores_rows_older_than_period():
    from datetime import datetime, timedelta, timezone

    with SessionLocal() as db:
        _seed(db, question="Старый вопрос", created_at=datetime.now(timezone.utc) - timedelta(days=100))
        _seed(db, question="Свежий вопрос")
        db.commit()
        report = gaps.build_report(db, days=30)
    assert report["summary"]["questions"] == 1 and report["gaps"][0]["topic"] == "Свежий вопрос"


def test_document_stats_flag_weak_and_unused_documents():
    src = '[{"relative_path": "A/плохой.md", "source_type": "context"}, {"relative_path": "A/связанный.md", "source_type": "related"}]'
    with SessionLocal() as db:
        for _ in range(2):
            _seed(db, question="q", kind="answer", rating=-1, sources_json=src)
        for i in range(20):  # достаточно вопросов, чтобы судить о неиспользуемых
            _seed(db, question=f"вопрос {i}", kind="answer", rating=1,
                  sources_json='[{"relative_path": "A/хороший.md", "source_type": "context"}]')
        db.commit()
        report = gaps.build_report(
            db,
            days=30,
            known_documents={"A/плохой.md": "плохой.md", "A/хороший.md": "хороший.md",
                             "A/связанный.md": "связанный.md", "A/забытый.md": "забытый.md"},
        )
    docs = report["documents"]
    assert [d["file_name"] for d in docs["weak"]] == ["плохой.md"]
    # «связанный» только предлагался как похожий — это не использование в ответе
    assert sorted(d["file_name"] for d in docs["unused"]) == ["забытый.md", "связанный.md"]


def test_unused_documents_not_reported_on_small_samples():
    with SessionLocal() as db:
        _seed(db, question="q", kind="answer")
        db.commit()
        docs = gaps.build_report(db, days=30, known_documents={"A/x.md": "x.md"})["documents"]
    assert docs["unused"] == [] and docs["unused_checked"] is False


# ------------------------------ API: права и выгрузка ------------------------------


def test_report_endpoint_is_restricted_to_admin_and_manager(client):
    user = register(client, "plain@test.local")
    assert client.get("/api/admin/knowledge-gaps", headers=user).status_code == 403
    assert client.get("/api/admin/knowledge-gaps/export", headers=user).status_code == 403
    assert client.get("/api/admin/knowledge-gaps").status_code == 401
    assert client.get("/api/admin/knowledge-gaps", headers=_admin(client)).status_code == 200


def test_manager_can_read_report(client):
    register(client, "boss@test.local")
    admin = _admin(client)
    users = client.get("/api/admin/users", headers=admin).json()
    boss = next(u for u in users if u["email"] == "boss@test.local")
    client.put(f"/api/admin/users/{boss['id']}/roles", headers=admin, json={"roles": ["user", "manager"]})
    boss_headers = auth_headers(client, "boss@test.local", "user-pass-123")
    assert client.get("/api/admin/knowledge-gaps", headers=boss_headers).status_code == 200


def test_report_validates_period(client):
    admin = _admin(client)
    assert client.get("/api/admin/knowledge-gaps?days=0", headers=admin).status_code == 422
    assert client.get("/api/admin/knowledge-gaps?days=9999", headers=admin).status_code == 422


def test_export_is_xlsx_and_neutralizes_formulas_in_questions(client):
    with SessionLocal() as db:
        _seed(db, question='=HYPERLINK("http://evil.example","клик")')
        db.commit()
    resp = client.get("/api/admin/knowledge-gaps/export", headers=_admin(client))
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == ["Сводка", "Пробелы", "Документы", "Конструктор"]
    cell = wb["Пробелы"]["A2"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    assert cell.quotePrefix is True
