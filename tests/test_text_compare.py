"""Сравнение версий документа по словам: PDF (текст и скан), Word. Все тексты выдуманные."""

from __future__ import annotations

import base64
import io

import pytest
from docx import Document

from app.services import chat_documents as cd
from app.services import text_compare as tc
from tests.test_spec_registry import FONT, needs_ocr

LETTER = [
    "Письмо № 12-345 от 03.09.2026",
    "Уважаемые коллеги! Сообщаем, что поставка насосов-дозаторов переносится на 14 календарных дней "
    "в связи с задержкой комплектующих у завода-изготовителя.",
    "Стоимость поставки остаётся прежней и составляет 480 000 рублей, включая НДС.",
    "Просим подтвердить согласие с новым сроком до 20 сентября.",
]


def _doc(text: str, name: str = "a.pdf", scanned: bool = False, conf: float = 96.0) -> tc.TextDoc:
    words = [tc.Word(t, tc.word_key(t), "стр. 1", conf if scanned else 100.0, scanned) for t in text.split()]
    return tc.TextDoc(name, words, pages=1, scanned_pages=[1] if scanned else [])


def _w(text: str, scanned: bool = True, conf: float = 96.0) -> tc.Word:
    return tc.Word(text, tc.word_key(text), "стр. 1", conf, scanned)


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for i, text in enumerate(paragraphs):
        doc.add_paragraph(text)
        if table and i == 0:
            t = doc.add_table(rows=len(table), cols=len(table[0]))
            for r, row in enumerate(table):
                for c, value in enumerate(row):
                    t.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _text_pdf(paragraphs: list[str]) -> bytes:
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page(width=595, height=842)
    y = 60
    for text in paragraphs:
        used = page.insert_textbox(pymupdf.Rect(60, y, 535, y + 150), text, fontsize=12, fontname="f", fontfile=FONT)
        y += 150 - used + 14
    page.insert_text((290, 815), "1", fontsize=9, fontname="f", fontfile=FONT)  # номер страницы
    return pdf.tobytes()


def _scan(pdf_bytes: bytes) -> bytes:
    import pymupdf
    from PIL import Image

    src, out = pymupdf.open(stream=pdf_bytes, filetype="pdf"), pymupdf.open()
    for page in src:
        pix = page.get_pixmap(dpi=200)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples).convert("L")
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=60)
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, stream=buf.getvalue())
    return out.tobytes()


needs_font = pytest.mark.skipif(FONT is None, reason="нужен шрифт с кириллицей")


# ------------------------------ слова и строки ------------------------------


@pytest.mark.parametrize("raw, key", [
    ("Договор", "договор"), ("«Ромашка»,", "ромашка"), ("1,5", "1.5"), ("1.5", "1.5"),
    ("Tpуба", "труба"),  # латинские T и p на скане
    ("№", ""), ("—", ""), ("12-345", "12345"),
])
def test_word_key(raw, key):
    assert tc.word_key(raw) == key


def test_hyphenated_words_are_joined():
    lines = [[("поставка", 100.0), ("дого-", 100.0)], [("вора", 90.0), ("продлена", 100.0)], [("Новый", 100.0)]]
    assert tc._join_hyphens(lines) == [[("поставка", 100.0), ("договора", 90.0)], [("продлена", 100.0)], [("Новый", 100.0)]]
    # перед заглавной — не перенос («Северо-Запад» на двух строках не склеиваем)
    assert tc._join_hyphens([[("Северо-", 100.0)], [("Запад", 100.0)]]) == [[("Северо-", 100.0)], [("Запад", 100.0)]]


def test_page_numbers_and_running_headers_are_dropped():
    def line(text):
        return [(w, 100.0) for w in text.split()]

    body = {
        1: ["Поставка переносится на две недели.", "Причина — задержка комплектующих.", "Срок уточним письмом."],
        2: ["Стоимость остаётся прежней.", "Оплата — по графику.", "Аванс не предусмотрен."],
        3: ["Ответ ждём до пятницы.", "Приложение: график поставки.", "Исполнитель: отдел снабжения."],
    }
    pages = {n: [line("ООО «Ромашка» Исх. № 7/2026"), *[line(t) for t in body[n]], line(f"Стр. {n} из 3")]
             for n in body}
    assert tc._drop_page_furniture(pages) == 6
    assert all([" ".join(w for w, _c in ln) for ln in pages[n]] == body[n] for n in body)


# ------------------------------- сравнение ------------------------------------


def test_text_layer_changes_are_reported_as_is():
    old = _doc(" ".join(LETTER), "letter_R01.pdf")
    new_text = " ".join(LETTER).replace("14 календарных", "21 календарный").replace(
        "до 20 сентября.", "до 20 сентября. Ответ направьте по электронной почте.")
    new = _doc(new_text.replace(" прежней", ""), "letter_R02.pdf")
    result = tc.compare(new, old)  # приложены в обратном порядке — старая по номеру редакции
    assert result.old.filename == "letter_R01.pdf" and result.same_document
    got = [(c.status, " ".join(w.text for w in c.old), " ".join(w.text for w in c.new)) for c in result.changes]
    assert got == [
        (tc.STATUS_CHANGED, "14 календарных", "21 календарный"),
        (tc.STATUS_REMOVED, "прежней", ""),
        (tc.STATUS_ADDED, "", "Ответ направьте по электронной почте."),
    ]


@pytest.mark.parametrize("old, new, status", [
    ([_w("480")], [_w("520")], tc.STATUS_CHANGED),  # число на число, прочитано уверенно
    ([_w("480", conf=40)], [_w("520")], tc.STATUS_CHECK),  # второе прочтение не сошлось
    ([_w("дотовор")], [_w("договор")], tc.STATUS_CHECK),  # почти то же написание
    ([_w("LUT")], [_w("шт")], tc.STATUS_CHECK),  # латиница против кириллицы
    ([_w("14"), _w("дней")], [_w("14")], tc.STATUS_CHECK),  # разная длина — пропуск распознавания
    ([_w("насосов")], [_w("компрессоров")], tc.STATUS_CHANGED),
    ([], [_w("Аванс"), _w("не")], tc.STATUS_CHECK),  # короткая вставка рядом со сканом
    ([], [_w(t) for t in "Аванс по договору не предусмотрен".split()], tc.STATUS_ADDED),
])
def test_scan_rules(old, new, status):
    assert tc.classify(old, new, scan=True)[0] == status


def test_text_layer_is_trusted_even_for_short_changes():
    assert tc.classify([_w("дотовор", scanned=False)], [_w("договор", scanned=False)], scan=False)[0] == tc.STATUS_CHANGED
    assert tc.classify([], [_w("не", scanned=False)], scan=False)[0] == tc.STATUS_ADDED


def test_unrelated_documents_are_not_compared():
    other = "Акт осмотра площадки: ограждение установлено, бытовые помещения размещены, подъездные пути " \
            "отсыпаны щебнем, освещение работает, замечаний к подрядчику нет, акт составлен в двух экземплярах"
    result = tc.compare(_doc(" ".join(LETTER)), _doc(other, "b.pdf"))
    assert not result.same_document and result.changes == [] and result.shared < 0.1
    forced = tc.compare(_doc(" ".join(LETTER)), _doc(other, "b.pdf"), force=True)
    assert forced.changes


def test_short_texts_compare_by_word_sets():
    assert tc.shared_share(["первый", "второй"], ["первый", "второй", "изменен"]) == 1.0


def test_numbers_reread_disagreeing_become_unsure(monkeypatch):
    import numpy as np
    import pytesseract

    gray = np.full((200, 400), 255, dtype=np.uint8)
    words = [["480", 95.0, (20, 20, 60, 30)], ["рублей", 95.0, (100, 20, 120, 30)], ["14", 95.0, (20, 100, 40, 30)]]

    def fake(image, **kw):  # первое число второе прочтение подтвердило, второе — нет
        return {"text": ["480", "11"], "top": [30, 30 + 106], "height": [60, 60]}

    monkeypatch.setattr(pytesseract, "image_to_data", fake)
    tc._reread_numbers(gray, words)
    assert [w[1] for w in words] == [95.0, 95.0, tc.UNSURE_CONF]


def test_word_keeps_tables_in_place():
    content = _docx(["Перед таблицей.", "После таблицы."], table=[["Поз.", "Наименование"], ["1", "Насос-дозатор"]])
    words = [w.text for w in tc.load_word(content, "a.docx").words]
    assert words == ["Перед", "таблицей.", "Поз.", "Наименование", "1", "Насос-дозатор", "После", "таблицы."]


def test_report_marks_removed_and_added_words():
    old = _doc(" ".join(LETTER))
    new = _doc(" ".join(LETTER).replace("480 000", "510 000"), "b.pdf")
    report = Document(io.BytesIO(tc.build_compare_docx(tc.compare(old, new))))
    table = report.tables[0]
    assert len(table.rows) == 2
    was, now = table.rows[1].cells[2].paragraphs[0].runs, table.rows[1].cells[3].paragraphs[0].runs
    assert any(r.font.strike and r.text.strip() == "480" for r in was)
    assert any(r.bold and r.text.strip() == "510" for r in now)
    assert "Как читать" in [p.text for p in report.paragraphs]


# ----------------------------------- чат ------------------------------------


@needs_font
def test_chat_compares_pdf_with_word():
    new = list(LETTER)
    new[2] = new[2].replace("480 000", "520 000")
    files = [("letter.docx", _docx(LETTER)), ("letter_signed.pdf", _text_pdf(new))]
    result = cd.handle_chat_files("старая — первая", files, llm=None)  # PDF и Word — сравнение по умолчанию
    assert "Изменено: **1**" in result.answer and "**480** 000" in result.answer and "**520** 000" in result.answer
    assert result.attachment.filename == "comparison.docx"


@needs_font
def test_chat_identical_documents_have_no_attachment():
    result = cd.handle_chat_files("сравни, старая — первая", [("a.docx", _docx(LETTER)), ("b.pdf", _text_pdf(LETTER))], llm=None)
    assert "Различий в тексте нет" in result.answer and result.attachment is None


def _changed_letter() -> list[str]:
    new = list(LETTER)
    new[2] = new[2].replace("480 000", "520 000")
    return new


def test_chat_asks_which_version_is_old_instead_of_guessing():
    files = [("a.docx", _docx(_changed_letter())), ("b.docx", _docx(LETTER))]
    asked = cd.handle_chat_files("сравни", files, llm=None)
    assert "Какая версия старая?" in asked.answer and asked.attachment is None
    assert [s["label"] for s in asked.suggestions] == ["Старая — a.docx", "Старая — b.docx"]

    # кнопка «Старая — b.docx» повторно отправляет те же файлы со своим текстом
    answered = cd.handle_chat_files(asked.suggestions[1]["message"], files, llm=None)
    assert "Старая: «b.docx», новая: «a.docx» (старую версию указали вы)" in answered.answer
    assert "**480** 000" in answered.answer and answered.attachment is not None


def test_chat_understands_new_version_named_by_user():
    files = [("a.docx", _docx(LETTER)), ("b.docx", _docx(_changed_letter()))]
    result = cd.handle_chat_files("сравни, новая — «a.docx»", files, llm=None)
    assert "Старая: «b.docx», новая: «a.docx» (старую версию указали вы)" in result.answer


def test_chat_same_file_names_ask_by_position():
    files = [("письмо.docx", _docx(LETTER)), ("письмо.docx", _docx(_changed_letter()))]
    asked = cd.handle_chat_files("сравни", files, llm=None)
    assert [s["label"] for s in asked.suggestions] == ["Старая — первый приложенный файл",
                                                       "Старая — второй приложенный файл"]
    answered = cd.handle_chat_files(asked.suggestions[0]["message"], files, llm=None)
    assert "старую версию указали вы" in answered.answer and "**480** 000" in answered.answer


def test_chat_revision_in_names_needs_no_question():
    files = [("письмо ред.2.docx", _docx(_changed_letter())), ("письмо ред.1.docx", _docx(LETTER))]
    result = cd.handle_chat_files("сравни", files, llm=None)
    assert result.suggestions == [] and "по номеру редакции" in result.answer


def test_answer_to_the_question_does_not_parse_files_again(monkeypatch):
    calls = []
    load = tc.load_document
    monkeypatch.setattr(tc, "load_document", lambda content, name: calls.append(name) or load(content, name))
    files = [("a.docx", _docx(_changed_letter())), ("b.docx", _docx(LETTER))]
    asked = cd.handle_chat_files("сравни", files, llm=None)
    cd.handle_chat_files(asked.suggestions[0]["message"], files, llm=None)
    assert sorted(calls) == ["a.docx", "b.docx"]  # разобраны один раз — при вопросе


def test_chat_refuses_unrelated_word_files():
    other = _docx(["Акт осмотра площадки: ограждение установлено, бытовые помещения размещены, подъездные пути "
                   "отсыпаны щебнем, освещение работает, замечаний к подрядчику нет."])
    result = cd.handle_chat_files("сравни", [("a.docx", _docx(LETTER)), ("b.docx", other)], llm=None)
    assert "не версии одного документа" in result.answer and result.attachment is None
    forced = cd.handle_chat_files("сравни всё равно, старая — первая", [("a.docx", _docx(LETTER)), ("b.docx", other)], llm=None)
    assert "по вашей просьбе" in forced.answer


@needs_ocr
@needs_font
def test_end_to_end_word_against_scan():
    new = list(LETTER)
    new[2] = new[2].replace("480 000", "520 000")
    files = [("letter.docx", _docx(LETTER)), ("letter_scan.pdf", _scan(_text_pdf(new)))]
    result = cd.handle_chat_files("сравни, старая — первая", files, llm=None)
    assert "Сканы распознаны автоматически" in result.answer
    assert "Изменено: **1**, добавлено: **0**, удалено: **0**" in result.answer
    assert "**480** 000" in result.answer and "**520** 000" in result.answer
    report = Document(io.BytesIO(base64.b64decode(result.attachment.content_base64)))
    assert len(report.tables[0].rows) == 2
