"""Тесты конвертации PDF → Word/Excel (Трек 1 плана конференции).

PDF генерируются на лету через PyMuPDF — внешние файлы не нужны. «Скан» —
страница с растровой картинкой без текстового слоя (именно её надо отдать в OCR).
OCR-тесты пропускаются, если в системе нет Tesseract (и русского языка).
"""

from __future__ import annotations

import base64
import io
import os
import shutil
import subprocess

import pymupdf
import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from app.services import chat_documents as cd
from app.services import document_processing as dp
from tests.conftest import ADMIN_EMAIL, ADMIN_PASSWORD, auth_headers


# ------------------------------ генерация PDF ----------------------------


def _text_page(doc, text: str, *, y: float = 60) -> "pymupdf.Page":
    page = doc.new_page()
    page.insert_htmlbox(pymupdf.Rect(72, y, 520, y + 60), f"<p>{text}</p>")
    return page


def _draw_table(page, rows: list[list[str]], *, y0: float = 150) -> None:
    x0, cw, rh = 72, 150, 30
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            rect = pymupdf.Rect(x0 + c * cw, y0 + r * rh, x0 + (c + 1) * cw, y0 + (r + 1) * rh)
            page.draw_rect(rect, color=(0, 0, 0), width=1)
            page.insert_htmlbox(rect + (3, 3, -3, -3), f"<span>{value}</span>")


def _text_pdf(*texts: str) -> bytes:
    doc = pymupdf.open()
    for t in texts:
        _text_page(doc, t)
    return doc.tobytes()


def _table_pdf(rows: list[list[str]], title: str = "Спецификация поставки материалов") -> bytes:
    doc = pymupdf.open()
    page = _text_page(doc, title)
    _draw_table(page, rows)
    return doc.tobytes()


_CYRILLIC_FONTS = (
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
)


def _find_font() -> str | None:
    return next((f for f in _CYRILLIC_FONTS if os.path.exists(f)), None)


def _scan_pdf(text: str, *, font_path: str | None = None) -> bytes:
    """PDF из одной страницы-картинки: текстового слоя нет."""
    font = ImageFont.truetype(font_path, 48) if font_path else ImageFont.load_default(size=48)
    img = Image.new("RGB", (1240, 400), "white")
    ImageDraw.Draw(img).text((60, 150), text, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")

    doc = pymupdf.open()
    page = doc.new_page(width=620, height=200)
    page.insert_image(page.rect, stream=buf.getvalue())
    return doc.tobytes()


def _has_tesseract_lang(lang: str) -> bool:
    exe = shutil.which("tesseract")
    if not exe:
        return False
    out = subprocess.run([exe, "--list-langs"], capture_output=True, text=True).stdout
    return lang in out.split()


needs_ocr = pytest.mark.skipif(not _has_tesseract_lang("eng"), reason="Tesseract не установлен")


class _StubLLM:
    def summarize_text(self, text, **kw):  # pragma: no cover — PDF LLM не зовёт
        raise AssertionError("LLM не должен вызываться при конвертации PDF")


def _decode_attachment(res: cd.ChatFileResult) -> bytes:
    assert res.attachment is not None
    return base64.b64decode(res.attachment.content_base64)


# ============================= parse_pdf ==================================


def test_parse_text_pdf_pages():
    data = dp.parse_pdf(_text_pdf("Договор поставки № 15", "Акт выполненных работ"), "doc.pdf")
    assert [p.page for p in data.pages] == [1, 2]
    assert "Договор поставки" in data.pages[0].text
    assert "Акт выполненных работ" in data.pages[1].text
    assert data.scanned_pages == []
    assert data.tables == []


def test_parse_pdf_extracts_table_and_excludes_it_from_text():
    rows = [["Наименование", "Кол-во"], ["Труба", "10"], ["Фланец", "4"]]
    data = dp.parse_pdf(_table_pdf(rows), "spec.pdf")
    assert len(data.tables) == 1
    assert data.tables[0].rows == rows
    # Заголовок остался текстом, а содержимое таблицы не задвоилось в тексте.
    assert "Спецификация" in data.full_text
    assert "Фланец" not in data.full_text


def test_parse_pdf_rejects_non_pdf_extension():
    with pytest.raises(dp.DocumentError):
        dp.parse_pdf(b"x", "doc.txt")


def test_parse_pdf_rejects_garbage():
    with pytest.raises(dp.DocumentError):
        dp.parse_pdf("это не pdf".encode(), "bad.pdf")


def test_parse_pdf_rejects_too_big(monkeypatch):
    monkeypatch.setattr(dp, "MAX_FILE_SIZE_PDF", 10)
    with pytest.raises(dp.DocumentError, match="слишком большой"):
        dp.parse_pdf(_text_pdf("текст"), "big.pdf")


def test_parse_pdf_rejects_too_many_pages(monkeypatch):
    monkeypatch.setattr(dp, "MAX_PDF_PAGES", 1)
    with pytest.raises(dp.DocumentError, match="слишком много страниц"):
        dp.parse_pdf(_text_pdf("Первая страница документа", "Вторая страница документа"), "long.pdf")


def test_parse_pdf_rejects_encrypted():
    doc = pymupdf.open()
    _text_page(doc, "Секретный документ")
    encrypted = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
    with pytest.raises(dp.DocumentError, match="паролем"):
        dp.parse_pdf(encrypted, "locked.pdf")


# ============================ сканы и OCR =================================


def test_scan_without_ocr_is_flagged(monkeypatch):
    """Нет Tesseract → страница помечена сканом, а ocr_available=False (без падения)."""
    monkeypatch.setattr(dp.shutil, "which", lambda name: None)
    data = dp.parse_pdf(_scan_pdf("Invoice 12345"), "scan.pdf")
    assert data.scanned_pages == [1]
    assert data.ocr_available is False
    assert data.full_text.strip() == ""


def test_chat_scan_without_ocr_gives_clear_error(monkeypatch):
    monkeypatch.setattr(dp.shutil, "which", lambda name: None)
    with pytest.raises(dp.DocumentError, match="OCR"):
        cd.handle_chat_files("переведи в word", [("scan.pdf", _scan_pdf("Invoice 12345"))], _StubLLM())


@needs_ocr
def test_scan_is_recognized_by_ocr():
    data = dp.parse_pdf(_scan_pdf("Invoice number 12345"), "scan.pdf")
    assert data.scanned_pages == [1]
    assert data.ocr_available is True
    assert "12345" in data.full_text


@pytest.mark.skipif(
    not (_has_tesseract_lang("rus") and _find_font()), reason="нет Tesseract rus или кириллического шрифта"
)
def test_russian_scan_is_recognized_by_ocr():
    data = dp.parse_pdf(_scan_pdf("Акт выполненных работ", font_path=_find_font()), "scan.pdf")
    assert "выполненных" in data.full_text.lower()


@needs_ocr
def test_mixed_pdf_marks_only_scanned_pages():
    doc = pymupdf.open()
    _text_page(doc, "Страница с текстовым слоем")
    scan = pymupdf.open(stream=_scan_pdf("Invoice number 777"), filetype="pdf")
    doc.insert_pdf(scan)
    data = dp.parse_pdf(doc.tobytes(), "mixed.pdf")
    assert data.scanned_pages == [2]
    assert "текстовым слоем" in data.pages[0].text


# ============================ сборка файлов ===============================


def test_build_docx_contains_text_and_table():
    rows = [["Наименование", "Кол-во"], ["Труба", "10"], ["Фланец", "4"]]
    data = dp.parse_pdf(_table_pdf(rows), "spec.pdf")
    docx = Document(io.BytesIO(dp.build_docx_from_pdf(data)))
    assert any("Спецификация" in p.text for p in docx.paragraphs)
    assert len(docx.tables) == 1
    assert [[c.text for c in r.cells] for r in docx.tables[0].rows] == rows


def _form_pdf() -> bytes:
    """Заголовок, таблица и текстовые поля ПОД таблицей — типичная форма-заявка."""
    doc = pymupdf.open()
    page = _text_page(doc, "Заявка на поставку материалов")
    _draw_table(page, [["Наименование", "Кол-во"], ["Труба", "10"], ["Фланец", "4"]])
    page.insert_htmlbox(pymupdf.Rect(72, 270, 520, 300), "<p>Ответственный: Иванов И.И.</p>")
    page.insert_htmlbox(pymupdf.Rect(72, 310, 520, 340), "<p>Срок поставки: 15.11.2026</p>")
    return doc.tobytes()


def _xlsx_cells(xlsx: bytes) -> list[list]:
    ws = load_workbook(io.BytesIO(xlsx)).active
    return [[c.value for c in r] for r in ws.iter_rows()]


def test_build_xlsx_keeps_table_rows():
    rows = [["Наименование", "Кол-во"], ["Труба", "10"]]
    data = dp.parse_pdf(_table_pdf(rows), "spec.pdf")
    wb = load_workbook(io.BytesIO(dp.build_xlsx_from_pdf(data)))
    assert wb.sheetnames == ["Стр1"]
    cells = _xlsx_cells(dp.build_xlsx_from_pdf(data))
    assert ["Наименование", "Кол-во"] in cells and ["Труба", "10"] in cells


def test_xlsx_keeps_text_fields_under_table_in_order():
    """Регрессия: в Excel пропадали текстовые поля вне таблицы."""
    cells = _xlsx_cells(dp.build_xlsx_from_pdf(dp.parse_pdf(_form_pdf(), "z.pdf")))
    flat = [row[0] for row in cells if row and row[0] is not None]
    for expected in ("Заявка на поставку материалов", "Ответственный: Иванов И.И.", "Срок поставки: 15.11.2026"):
        assert expected in flat
    # порядок как в PDF: заголовок → таблица → поля под ней
    assert flat.index("Заявка на поставку материалов") < flat.index("Наименование")
    assert flat.index("Фланец") < flat.index("Ответственный: Иванов И.И.") < flat.index("Срок поставки: 15.11.2026")


def test_docx_keeps_text_and_table_in_page_order():
    docx = Document(io.BytesIO(dp.build_docx_from_pdf(dp.parse_pdf(_form_pdf(), "z.pdf"))))
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    order = []
    for child in docx.element.body:
        if child.tag.endswith("}p"):
            text = Paragraph(child, docx).text.strip()
            if text:
                order.append(text)
        elif child.tag.endswith("}tbl"):
            assert Table(child, docx).rows
            order.append("<таблица>")
    assert order == [
        "Заявка на поставку материалов",
        "<таблица>",
        "Ответственный: Иванов И.И.",
        "Срок поставки: 15.11.2026",
    ]


def test_form_field_values_are_not_lost():
    doc = pymupdf.open()
    page = _text_page(doc, "Анкета сотрудника")
    widget = pymupdf.Widget()
    widget.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    widget.field_name = "Телефон"
    widget.rect = pymupdf.Rect(72, 200, 300, 220)
    widget.field_value = "+7 900 123-45-67"
    page.add_widget(widget)
    data = dp.parse_pdf(doc.tobytes(), "anketa.pdf")
    assert "+7 900 123-45-67" in data.full_text


def test_build_xlsx_without_tables_is_error():
    data = dp.parse_pdf(_text_pdf("Просто текст без таблиц"), "plain.pdf")
    with pytest.raises(dp.DocumentError, match="не найдено таблиц"):
        dp.build_xlsx_from_pdf(data)


# ============================ маршрутизация чата ==========================


def test_chat_pdf_defaults_to_word():
    res = cd.handle_chat_files("", [("letter.pdf", _text_pdf("Входящее письмо № 5"))], _StubLLM())
    assert res.attachment.filename == "letter.docx"
    assert res.attachment.mime.endswith("wordprocessingml.document")
    docx = Document(io.BytesIO(_decode_attachment(res)))
    assert any("Входящее письмо" in p.text for p in docx.paragraphs)
    assert "OCR" not in res.answer


def test_chat_pdf_excel_keyword_gives_xlsx():
    pdf = _table_pdf([["Наименование", "Кол-во"], ["Труба", "10"]])
    res = cd.handle_chat_files("переведи в Excel", [("spec.pdf", pdf)], _StubLLM())
    assert res.attachment.filename == "spec.xlsx"
    assert ["Труба", "10"] in _xlsx_cells(_decode_attachment(res))


@needs_ocr
def test_chat_scan_warns_about_ocr():
    res = cd.handle_chat_files("", [("scan.pdf", _scan_pdf("Invoice number 12345"))], _StubLLM())
    assert "OCR" in res.answer
    assert "Проверьте результат" in res.answer


def test_chat_pdf_conversion_of_several_files_is_rejected():
    pdf = _text_pdf("Документ")
    with pytest.raises(dp.DocumentError, match="по одному PDF"):
        cd.handle_chat_files("переведи в word", [("a.pdf", pdf), ("b.pdf", pdf)], _StubLLM())


def test_chat_pdf_cannot_mix_with_other_types():
    with pytest.raises(dp.DocumentError, match="одного типа"):
        cd.handle_chat_files("", [("a.pdf", _text_pdf("Документ")), ("b.xlsx", b"x")], _StubLLM())


def test_endpoint_converts_pdf(client):
    headers = auth_headers(client, ADMIN_EMAIL, ADMIN_PASSWORD)
    resp = client.post(
        "/api/ask-files",
        data={"message": "в word"},
        files={"files": ("letter.pdf", _text_pdf("Входящее письмо № 5"), "application/pdf")},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["attachment"]["filename"] == "letter.docx"
