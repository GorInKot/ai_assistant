"""Сравнение редакций спецификации по позициям: сопоставление, статусы, Excel, чат.

Все наименования, шифры и числа здесь выдуманные."""

from __future__ import annotations

import base64
import io

import pytest
from openpyxl import load_workbook

from app.services import chat_documents as cd
from app.services import spec_compare as sc
from app.services import spec_registry as sr
from app.services.document_processing import DocumentError

CODE = "12345678.001-Р-100.000.000-ТХ-01-С-001"


def _pos(name, qty="", unit="шт.", pos="", page=1, mass="", type_="", issues=()):
    p = sr.Position(page=page)
    p.cells = [pos, name, type_, "", "", unit, qty, mass, ""]
    p.issues = list(issues)
    return p


def _doc(filename, positions, code=CODE):
    pages = max((p.page for p in positions), default=0)
    return sr.SpecDocument(filename=filename, code=code, pages=pages, table_pages=pages, positions=positions)


def _base():
    return [
        _pos("Насос-дозатор плунжерный НД-25", "2", mass="48"),
        _pos("Регулятор давления прямого действия Ду50", "6", mass="7,5"),
        _pos("Труба стальная электросварная 114х5", "120", unit="м", mass="13,44"),
        _pos("Отвод крутоизогнутый 90-114х5", "14", mass="1,9"),
        _pos("Фильтр-грязевик Ду80", "1", mass="22"),
    ]


def _statuses(result):
    return [r.status for r in result.rows]


# ------------------------------ сопоставление ------------------------------


def test_same_revision_has_no_changes():
    result = sc.compare(_doc("a_R01.pdf", _base()), _doc("a_R02.pdf", _base()))
    assert _statuses(result) == [sc.STATUS_SAME] * 5
    assert result.same_document


def test_quantity_change_read_confidently_is_changed():
    new = _base()
    new[2] = _pos("Труба стальная электросварная 114х5", "95,5", unit="м", mass="13,44")
    result = sc.compare(_doc("a_R01.pdf", _base()), _doc("a_R02.pdf", new))
    row = result.rows[2]
    assert row.status == sc.STATUS_CHANGED and row.changes == ["кол-во: 120 → 95,5"]
    assert result.count(sc.STATUS_SAME) == 4


def test_change_with_unsure_reading_goes_to_check():
    new = _base()
    new[1] = _pos("Регулятор давления прямого действия Ду50", "8", mass="7,5", issues=["посторонние пометки у чисел (штамп, рука, облако правки)"])
    row = sc.compare(_doc("a_R01.pdf", _base()), _doc("a_R02.pdf", new)).rows[1]
    assert row.status == sc.STATUS_CHECK
    assert row.changes == ["кол-во: 6 → 8"] and row.reasons[0].startswith("нов.: посторонние пометки")


def test_ocr_noise_in_names_still_matches():
    # «н» прочитана как «и», кириллическая «о» — как латинская
    new = [_pos(p.cells[sr.COL_NAME].replace("н", "и").replace("о", "o"), p.cells[sr.COL_QTY],
                unit=p.cells[sr.COL_UNIT], mass=p.cells[sr.COL_MASS]) for p in _base()]
    result = sc.compare(_doc("a_R01.pdf", _base()), _doc("a_R02.pdf", new))
    assert _statuses(result) == [sc.STATUS_SAME] * 5


def test_added_and_removed_positions():
    new = _base()[:2] + [_pos("Розетка штепсельная двухполюсная", "3", mass="31")] + _base()[3:]
    result = sc.compare(_doc("a_R01.pdf", _base()), _doc("a_R02.pdf", new))
    assert result.count(sc.STATUS_ADDED) == 1 and result.count(sc.STATUS_REMOVED) == 1
    added = next(r for r in result.rows if r.status == sc.STATUS_ADDED)
    assert added.new.cells[sr.COL_NAME] == "Розетка штепсельная двухполюсная"


def test_unsure_position_without_pair_is_check_not_removed():
    old = _base() + [_pos("", ",", unit="Иш |", issues=["нет наименования"])]  # обрывок правки от руки
    result = sc.compare(_doc("a_R01.pdf", old), _doc("a_R02.pdf", _base()))
    assert result.count(sc.STATUS_REMOVED) == 0
    lonely = result.rows[-1]
    assert lonely.status == sc.STATUS_CHECK and lonely.changes == ["есть только в старой редакции"]


def test_position_number_and_quantity_match_despite_different_names():
    # вариант прочитан в одной редакции полным наименованием, в другой — обрывком
    old = [_pos("Провод монтажный многожильный, сечение, мм: 3х2,5", "1,25", unit="км", pos="2.1.4"),
           _pos("Провод монтажный многожильный, сечение, мм: 3х4", "0,8", unit="км", pos="2.1.5")]
    new = [_pos("Ax2,5", "1,25", unit="км", pos="2.1.4"), _pos("3х4", "0,8", unit="км", pos="2.1.5")]
    result = sc.compare(_doc("a_R01.pdf", old), _doc("a_R02.pdf", new))
    assert _statuses(result) == [sc.STATUS_SAME, sc.STATUS_SAME]


def test_dimension_change_in_name_is_check():
    new = _base()
    new[3] = _pos("Отвод крутоизогнутый 90-133х5", "14", mass="1,9")
    row = sc.compare(_doc("a_R01.pdf", _base()), _doc("a_R02.pdf", new)).rows[3]
    assert row.status == sc.STATUS_CHECK
    assert row.changes == ["размер в наименовании: 114х5 → 133х5"]


def test_numbers_and_units_compare_by_value():
    old = [_pos("Труба стальная 76х4", "12,0", unit="м", mass="27,1*"), _pos("Скоба крепёжная", "40", unit="шт")]
    new = [_pos("Труба стальная 76х4", "12", unit="м", mass="27.1"), _pos("Скоба крепёжная", "40", unit="шт.")]
    assert _statuses(sc.compare(_doc("a_R01.pdf", old), _doc("a_R02.pdf", new))) == [sc.STATUS_SAME] * 2


def test_moved_position_is_matched():
    new = _base()
    new = [new[4]] + new[:4]  # фильтр перенесён в начало
    result = sc.compare(_doc("a_R01.pdf", _base()), _doc("a_R02.pdf", new))
    assert result.count(sc.STATUS_SAME) == 5
    assert [r.moved for r in result.rows].count(True) == 1


def test_sheets_missing_in_partial_revision_are_not_added():
    # изменение выпущено только на лист 1; в полной редакции есть ещё лист 2
    full = _base() + [_pos("Кронштейн настенный", "16", page=2), _pos("Хомут червячный", "40", page=2)]
    result = sc.compare(_doc("a изм.1.pdf", _base()), _doc("a изм.2.pdf", full))
    assert result.count(sc.STATUS_ADDED) == 0
    assert result.count(sc.STATUS_NEW_SHEET) == 2 and result.pages(sc.STATUS_NEW_SHEET) == [2]


def test_excluded_positions_do_not_take_part():
    old = _base()
    old[0].excluded = True
    result = sc.compare(_doc("a_R01.pdf", old), _doc("a_R02.pdf", _base()[1:]))
    assert _statuses(result) == [sc.STATUS_SAME] * 4 and result.excluded == 1


# -------------------------- один ли это документ --------------------------


def test_different_documents_are_recognized():
    other = [_pos("Прожектор заливающего света", "48"), _pos("Выключатель автоматический 16А", "12"),
             _pos("Кабель-канал 40х25", "90", unit="м")]
    result = sc.compare(_doc("a.pdf", _base()), _doc("b.pdf", other, code="12345678.001-Р-100.000.000-ЭО-02-С-001"))
    assert not result.same_document
    assert "сопоставлено 0" in result.evidence


def test_same_document_by_content_when_codes_unread():
    new = _base()
    new[0] = _pos("Насос-дозатор плунжерный НД-25", "3", mass="48")
    result = sc.compare(_doc("a.pdf", _base(), code="a"), _doc("b.pdf", new, code="b"))
    assert result.same_document and result.evidence.startswith("по содержимому")


def test_same_quantities_required_for_similar_lists():
    # похожие перечни (тот же набор позиций), но количества другие — это другой документ
    other = [_pos(p.cells[sr.COL_NAME], str(int(float(p.cells[sr.COL_QTY].replace(",", "."))) + 7),
                  unit=p.cells[sr.COL_UNIT]) for p in _base()]
    result = sc.compare(_doc("a.pdf", _base(), code="a"), _doc("b.pdf", other, code="b"))
    assert not result.same_document


@pytest.mark.parametrize("first, second, old", [
    ("x_R04.pdf", "x_R03.pdf", "x_R03.pdf"),
    ("x изм.2.pdf", "x изм.1.pdf", "x изм.1.pdf"),
    ("x-rC02-f01.pdf", "x-rC01-f01.pdf", "x-rC01-f01.pdf"),
    ("x изм.3.pdf", "x_R04.pdf", "x изм.3.pdf"),  # разные схемы нумерации — по порядку файлов
])
def test_order_by_revision_in_filename(first, second, old):
    a, b = _doc(first, _base()), _doc(second, _base())
    assert sc.order_documents(a, b)[0].filename == old


# ---------------------------------- Excel -----------------------------------


def test_xlsx_statuses_colors_and_formula_safety():
    new = _base()
    new[2] = _pos("Труба стальная электросварная 114х5", "95,5", unit="м", mass="13,44")
    new.append(_pos("=HYPERLINK(\"x\")", "1"))
    result = sc.compare(_doc("a_R01.pdf", _base()), _doc("a_R02.pdf", new))
    wb = load_workbook(io.BytesIO(sc.build_compare_xlsx(result)))
    assert wb.sheetnames == ["Сравнение", "Сводка", "Как читать"]
    ws = wb["Сравнение"]
    assert [c.value for c in ws[1]] == sc.EXCEL_HEADERS
    changed = next(r for r in ws.iter_rows(min_row=2) if r[0].value == sc.STATUS_CHANGED)
    assert changed[9].value == 120 and changed[10].value == 95.5 and changed[11].value == -24.5
    assert changed[0].fill.fgColor.rgb.endswith(sc.FILLS[sc.STATUS_CHANGED])
    formula = next(r for r in ws.iter_rows(min_row=2) if r[0].value == sc.STATUS_ADDED)[5]
    assert formula.data_type == "s" and formula.quotePrefix


# ----------------------------------- чат ------------------------------------


def _fake_process(docs):
    def process(content, name, budget):
        return docs[name]
    return process


@pytest.fixture
def ocr_ready(monkeypatch):
    monkeypatch.setattr(sr, "unavailable_reason", lambda: None)


def test_chat_compares_two_spec_pdfs(monkeypatch, ocr_ready):
    new = _base()
    new[2] = _pos("Труба стальная электросварная 114х5", "95,5", unit="м", mass="13,44")
    docs = {"a_R02.pdf": _doc("a_R02.pdf", new), "a_R01.pdf": _doc("a_R01.pdf", _base())}
    monkeypatch.setattr(sr, "process_pdf", _fake_process(docs))
    result = cd.handle_chat_files("сравни редакции", [("a_R02.pdf", b"%PDF"), ("a_R01.pdf", b"%PDF")], llm=None)
    assert result.attachment.filename == "spec_comparison.xlsx"
    assert "Старая: «a_R01.pdf»" in result.answer and "Изменено: **1**" in result.answer
    assert "кол-во: 120 → 95,5" in result.answer


def test_chat_refuses_to_compare_different_documents(monkeypatch, ocr_ready):
    other = [_pos("Прожектор заливающего света", "48"), _pos("Выключатель автоматический 16А", "12")]
    docs = {"a.pdf": _doc("a.pdf", _base()), "b.pdf": _doc("b.pdf", other, code="12345678.001-Р-100.000.000-ЭО-02-С-001")}
    monkeypatch.setattr(sr, "process_pdf", _fake_process(docs))
    files = [("a.pdf", b"%PDF"), ("b.pdf", b"%PDF")]
    result = cd.handle_chat_files("что изменилось?", files, llm=None)
    assert result.attachment is None and "не версии одного документа" in result.answer
    forced = cd.handle_chat_files("сравни всё равно", files, llm=None)
    assert forced.attachment is not None and "по вашей просьбе" in forced.answer


def test_page_ranges():
    assert cd._page_ranges([2, 4, 5, 6, 9, 10]) == "2, 4–6, 9–10"
    assert cd._page_ranges([7]) == "7"


def test_chat_compare_needs_spec_form(monkeypatch, ocr_ready):
    docs = {"a.pdf": _doc("a.pdf", _base()), "letter.pdf": sr.SpecDocument(filename="letter.pdf", pages=2)}
    monkeypatch.setattr(sr, "process_pdf", _fake_process(docs))
    with pytest.raises(DocumentError, match="только спецификации"):
        cd.handle_chat_files("сравни", [("a.pdf", b"%PDF"), ("letter.pdf", b"%PDF")], llm=None)


# ------------------------- сквозной прогон по «сканам» -------------------------

from tests.test_spec_registry import ROWS, _synthetic_scan, needs_ocr  # noqa: E402


@needs_ocr
def test_end_to_end_two_synthetic_revisions():
    changed = [list(r) for r in ROWS]
    changed[1][4] = "16,0"  # количество трубы: 12,5 → 16,0
    files = [("spec_R01.pdf", _synthetic_scan()), ("spec_R02.pdf", _synthetic_scan(content=[tuple(r) for r in changed]))]
    result = cd.handle_chat_files("сравни спецификации", files, llm=None)
    assert "Изменено: **1**" in result.answer and "кол-во: 12,5 → 16,0" in result.answer
    ws = load_workbook(io.BytesIO(base64.b64decode(result.attachment.content_base64)))["Сравнение"]
    statuses = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    assert statuses.count(sc.STATUS_CHANGED) == 1 and statuses.count(sc.STATUS_SAME) == 3
