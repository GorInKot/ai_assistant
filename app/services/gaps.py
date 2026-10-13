"""Аналитика пробелов в знаниях: чего сотрудникам не хватает в базе документов.

Откуда данные: каждый вопрос к базе знаний пишется в журнал `qa_log` (что спросили, нашёлся
ли ответ, сколько утверждений без опоры в документах, какие документы показаны, оценил ли
пользователь ответ «полезно / не полезно»). Из журнала собирается отчёт для владельца базы:

- «В базе нет ответа» — вопросы, на которые ассистент ответил отказом или просьбой уточнить,
  после которой ответа так и не последовало;
- «Документ есть, но не помогает» — вопросы с ответом, который оценили как неполезный;
- документы, к которым чаще всего предъявляют претензии, и документы, которые не используются.

Похожие вопросы объединяются в темы (по эмбеддингам, а без них — по общим словам), чтобы
«как оформить загранкомандировку» и «командировка за границу» считались одной потребностью.
Считаем разных людей, а не обращения: один сотрудник, переформулировавший вопрос трижды,
— это один человек, а не три.

Ограничение: оценка и отказ — косвенные сигналы. Отчёт показывает, где искать проблему,
но решает, что дописать в базу, человек.
"""

from __future__ import annotations

import json
import logging
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any

import numpy as np
from sqlalchemy.orm import Session

from app.db import Message, QaLog, User
from app.services.document_processing import append_row
from app.services.grounding import STEM_LEN, content_stems

logger = logging.getLogger(__name__)

RATING_USEFUL = 1
RATING_NOT_USEFUL = -1

KIND_ANSWER = "answer"
KIND_MISSING = "missing"
KIND_CLARIFY = "clarify"
KIND_SMALLTALK = "smalltalk"

# Исходы, при которых ответа по документам пользователь не получил.
NO_ANSWER_KINDS = (KIND_MISSING, KIND_CLARIFY)
# Следующая реплика пользователя в течение этого срока считается ответом на просьбу уточнить
# (так же работает dialog_state: незавершённый вопрос живёт 20 минут и склеивается со следующим).
CLARIFY_WINDOW = timedelta(minutes=20)

# Когда два вопроса считаются одной темой. Пороги подобраны по замеру на qwen3-embedding:0.6b:
# короткие шаблонные вопросы («Как оформить командировку?» / «Как оформить отпуск?») имеют
# близость 0.74–0.77, хотя потребности разные, а настоящие пересказы — 0.6–0.9. Поэтому одной
# близости эмбеддингов мало: сначала сравниваются «предметные» слова вопроса (без «как»,
# «оформить», «порядок»…), эмбеддинг только подтверждает или ловит пересказ другими словами.
# Лучше не объединить два похожих вопроса, чем слить разные темы в одну.
TOPIC_JACCARD = 0.5
EMBEDDING_WITH_SHARED_TERM = 0.8  # и есть хотя бы одно общее предметное слово
EMBEDDING_ALONE = 0.88  # пересказ совсем другими словами
MAX_QUESTIONS_TO_CLUSTER = 500  # самые свежие; ограничивает и время запроса к эмбеддеру

MAX_EXAMPLES = 4
MAX_TOP_SOURCES = 3
# Документ «проблемный», если его ответы оценили как неполезные не менее стольких раз
# и не реже, чем как полезные.
WEAK_DOC_MIN_DOWN = 2
# Выводы «документ не используется» осмысленны только при достаточном числе вопросов.
MIN_QUESTIONS_FOR_UNUSED = 20


# ------------------------------ запись в журнал ------------------------------


def record_answer(
    db: Session,
    *,
    user: User,
    response: Any,
    fallback_question: str,
    message_id: int | None,
) -> None:
    """Записать ответ в журнал вопросов. Ошибка журнала не должна ломать ответ пользователю."""
    if response.kind == KIND_SMALLTALK:
        return
    try:
        grounding = response.grounding or {}
        sources = [
            {"relative_path": s.get("relative_path"), "source_type": s.get("source_type", "context")}
            for s in (response.log_sources if response.log_sources is not None else response.sources)
            if s.get("relative_path")
        ]
        db.add(
            QaLog(
                message_id=message_id,
                user_id=user.id,
                question=(response.resolved_question or fallback_question).strip(),
                kind=response.kind,
                division=user.division,
                claims_total=grounding.get("total"),
                claims_unsupported=len(grounding.get("unsupported", [])) if grounding else None,
                sources_json=json.dumps(sources, ensure_ascii=False),
            )
        )
        db.commit()
    except Exception:  # noqa: BLE001 — журнал вторичен
        db.rollback()
        logger.exception("Не удалось записать вопрос в журнал")


def set_rating(db: Session, message: Message, rating: int | None) -> None:
    """Сохранить оценку ответа (в сообщении и в журнале вопросов)."""
    if rating not in (RATING_USEFUL, RATING_NOT_USEFUL, None):
        raise ValueError("rating должен быть 1, -1 или null")
    message.rating = rating
    for row in db.query(QaLog).filter(QaLog.message_id == message.id).all():
        row.rating = rating
    db.commit()


# ------------------------------ объединение вопросов в темы ------------------------------

# Слова-шаблоны вопроса: они не говорят, о чём спрашивают, и только сближают разные вопросы.
_GENERIC_WORDS = """
как какой какая какие каких каком какую где куда откуда когда кто что чего кому почему зачем
сколько ли можно нужно надо необходимо должен порядок процедура процесс правила инструкция
алгоритм шаги этапы оформить оформление оформлять заказать получить получение подать подача
сделать узнать расскажите подскажите объясните помогите хочу хотел требуется нужна нужен нужны
обращаться обратиться пройти
""".split()
_GENERIC_STEMS = frozenset(w[:STEM_LEN] for w in _GENERIC_WORDS)


def topic_terms(question: str) -> frozenset[str]:
    """Предметные слова вопроса: основы значимых слов без слов-шаблонов («как», «оформить»…)."""
    stems = content_stems(question)
    return frozenset(stems - _GENERIC_STEMS) or frozenset(stems)


def _same_topic(a: frozenset[str], b: frozenset[str], similarity: float | None) -> bool:
    shared = len(a & b)
    if shared and shared / len(a | b) >= TOPIC_JACCARD:
        return True
    if similarity is None:
        return False
    return similarity >= EMBEDDING_ALONE or (shared > 0 and similarity >= EMBEDDING_WITH_SHARED_TERM)


def _embed(questions: list[str], embedder: Any) -> np.ndarray | None:
    """Нормированные векторы вопросов или None, если эмбеддер недоступен."""
    if embedder is None or not questions:
        return None
    try:
        vectors = np.asarray(embedder.embed_documents(questions), dtype=np.float32)
    except Exception:  # noqa: BLE001 — сеть/модель недоступны → запасной способ
        logger.warning("Эмбеддер недоступен, вопросы объединяются по общим словам", exc_info=True)
        return None
    norms = np.linalg.norm(vectors, axis=1, keepdims=True)
    norms[norms == 0] = 1.0
    return vectors / norms


def cluster_questions(questions: list[str], embedder: Any = None) -> tuple[list[list[int]], bool]:
    """Сгруппировать вопросы по теме. Возвращает (группы из индексов, использованы ли эмбеддинги).

    Вопросы связываются попарно (см. `_same_topic`), тема — компонента связности. Порядок
    внутри группы и порядок групп сохраняют порядок входа: на входе свежие вопросы первыми,
    поэтому тему называет самый свежий вопрос.
    """
    n = len(questions)
    vectors = _embed(questions, embedder)
    similarity = vectors @ vectors.T if vectors is not None else None
    terms = [topic_terms(q) for q in questions]

    parent = list(range(n))

    def find(x: int) -> int:
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    for i in range(n):
        for j in range(i + 1, n):
            if find(i) != find(j) and _same_topic(
                terms[i], terms[j], float(similarity[i, j]) if similarity is not None else None
            ):
                parent[find(j)] = find(i)

    groups: dict[int, list[int]] = {}
    for i in range(n):
        groups.setdefault(find(i), []).append(i)
    return sorted(groups.values(), key=lambda g: g[0]), vectors is not None


# ------------------------------ отчёт ------------------------------


@dataclass
class Gap:
    topic: str  # формулировка темы — типичный вопрос
    problem: str  # no_document | weak_document
    requests: int
    people: int
    divisions: list[str]
    not_useful: int
    unanswered: int  # отказов и просьб уточнить без ответа
    warned: int
    examples: list[str]
    sources: list[dict]
    last_asked: str | None = None

    def to_dict(self) -> dict:
        return self.__dict__.copy()


PROBLEM_NO_DOCUMENT = "no_document"
PROBLEM_WEAK_DOCUMENT = "weak_document"
PROBLEM_LABELS = {
    PROBLEM_NO_DOCUMENT: "В базе нет ответа",
    PROBLEM_WEAK_DOCUMENT: "Документ есть, но ответ не помогает",
}


def _dead_end_clarifications(rows: list[QaLog]) -> set[int]:
    """Просьбы уточнить, после которых ответа не было: id строк журнала.

    Следующая реплика того же сотрудника в течение 20 минут — это и есть уточнение (диалог
    склеивает её с вопросом). Если она получила ответ — вопрос решён, а если нет (или её
    не было) — человек остался без ответа.
    """
    by_user: dict[int | None, list[QaLog]] = {}
    for row in sorted(rows, key=lambda r: r.id):
        by_user.setdefault(row.user_id, []).append(row)

    dead: set[int] = set()
    for user_rows in by_user.values():
        for i, row in enumerate(user_rows):
            if row.kind != KIND_CLARIFY:
                continue
            nxt = user_rows[i + 1] if i + 1 < len(user_rows) else None
            resolved = (
                nxt is not None
                and nxt.kind == KIND_ANSWER
                and (nxt.created_at - row.created_at) <= CLARIFY_WINDOW
            )
            if not resolved:
                dead.add(row.id)
    return dead


def _is_problem(row: QaLog, dead_ends: set[int]) -> bool:
    """Строка журнала, которая указывает на пробел.

    Отказ — всегда пробел, даже если пользователь поставил «полезно»: так отмечают честный
    ответ «в базе нет», а потребность остаётся неудовлетворённой. Ответ или просьба уточнить
    — только если оценены как «не полезно» (просьба — ещё и если без ответа заглохла).
    """
    if row.kind == KIND_MISSING:
        return True
    if row.rating == RATING_NOT_USEFUL:
        return row.kind in (KIND_ANSWER, KIND_CLARIFY)
    return row.kind == KIND_CLARIFY and row.id in dead_ends


def _context_paths(row: QaLog) -> list[str]:
    try:
        items = json.loads(row.sources_json or "[]")
    except json.JSONDecodeError:
        return []
    return [i["relative_path"] for i in items if i.get("source_type", "context") == "context"]


def _iso(value: datetime | None) -> str | None:
    return value.isoformat() if value else None


def _load_rows(db: Session, days: int) -> list[QaLog]:
    cutoff = datetime.now(timezone.utc) - timedelta(days=days)
    return (
        db.query(QaLog)
        .filter(QaLog.created_at >= cutoff)
        .order_by(QaLog.created_at.desc(), QaLog.id.desc())
        .all()
    )


def _build_gaps(rows: list[QaLog], embedder: Any) -> tuple[list[Gap], bool]:
    dead_ends = _dead_end_clarifications(rows)
    problems = [r for r in rows if _is_problem(r, dead_ends)][:MAX_QUESTIONS_TO_CLUSTER]
    clusters, used_embeddings = cluster_questions([r.question for r in problems], embedder)

    gaps: list[Gap] = []
    for members in clusters:
        group = [problems[i] for i in members]
        unanswered = sum(1 for r in group if r.kind in NO_ANSWER_KINDS)
        sources: dict[str, int] = {}
        for r in group:
            for path in _context_paths(r):
                sources[path] = sources.get(path, 0) + 1
        examples: list[str] = []
        for r in group:
            if r.question not in examples:
                examples.append(r.question)
        gaps.append(
            Gap(
                topic=group[0].question,
                problem=PROBLEM_NO_DOCUMENT if unanswered * 2 >= len(group) else PROBLEM_WEAK_DOCUMENT,
                requests=len(group),
                people=len({r.user_id for r in group if r.user_id is not None}) or len(group),
                divisions=sorted({r.division for r in group if r.division}),
                not_useful=sum(1 for r in group if r.rating == RATING_NOT_USEFUL),
                unanswered=unanswered,
                warned=sum(1 for r in group if (r.claims_unsupported or 0) > 0),
                examples=examples[:MAX_EXAMPLES],
                sources=[
                    {"relative_path": p, "count": n}
                    for p, n in sorted(sources.items(), key=lambda kv: -kv[1])[:MAX_TOP_SOURCES]
                ],
                last_asked=_iso(group[0].created_at),
            )
        )
    gaps.sort(key=lambda g: (-g.people, -g.requests, g.topic))
    return gaps, used_embeddings


def _document_stats(rows: list[QaLog], known_documents: dict[str, str]) -> dict:
    stats: dict[str, dict[str, int]] = {}
    for r in rows:
        if r.kind != KIND_ANSWER:
            continue
        for path in set(_context_paths(r)):
            s = stats.setdefault(path, {"used": 0, "useful": 0, "not_useful": 0, "warned": 0})
            s["used"] += 1
            s["useful"] += r.rating == RATING_USEFUL
            s["not_useful"] += r.rating == RATING_NOT_USEFUL
            s["warned"] += (r.claims_unsupported or 0) > 0

    def name(path: str) -> str:
        return known_documents.get(path) or path.rsplit("/", 1)[-1]

    weak = [
        {"relative_path": p, "file_name": name(p), **s}
        for p, s in stats.items()
        if s["not_useful"] >= WEAK_DOC_MIN_DOWN and s["not_useful"] >= s["useful"]
    ]
    weak.sort(key=lambda d: (-d["not_useful"], d["file_name"]))

    questions = sum(1 for r in rows if r.kind in (KIND_ANSWER, KIND_MISSING))
    unused: list[dict] = []
    if questions >= MIN_QUESTIONS_FOR_UNUSED:
        unused = [
            {"relative_path": p, "file_name": n}
            for p, n in sorted(known_documents.items(), key=lambda kv: kv[1])
            if p not in stats
        ]
    return {"weak": weak, "unused": unused, "unused_checked": questions >= MIN_QUESTIONS_FOR_UNUSED}


def build_report(
    db: Session,
    *,
    days: int = 30,
    embedder: Any = None,
    known_documents: dict[str, str] | None = None,
) -> dict:
    """Отчёт о пробелах в знаниях за последние `days` дней.

    `known_documents` — {относительный путь: имя файла} документов, участвующих в поиске;
    нужен, чтобы найти документы, которые ни разу не использовались.
    """
    rows = _load_rows(db, days)
    kb_rows = [r for r in rows if r.kind in (KIND_ANSWER, KIND_MISSING, KIND_CLARIFY)]
    rated = [r for r in kb_rows if r.rating is not None]
    useful = sum(1 for r in rated if r.rating == RATING_USEFUL)
    gaps, used_embeddings = _build_gaps(rows, embedder)

    summary = {
        # Каждая реплика пользователя, дошедшая до базы знаний (включая просьбы уточнить).
        "questions": len(kb_rows),
        "people": len({r.user_id for r in kb_rows if r.user_id is not None}),
        "answered": sum(1 for r in kb_rows if r.kind == KIND_ANSWER),
        "refused": sum(1 for r in kb_rows if r.kind == KIND_MISSING),
        "clarifications": sum(1 for r in kb_rows if r.kind == KIND_CLARIFY),
        "warned": sum(1 for r in kb_rows if r.kind == KIND_ANSWER and (r.claims_unsupported or 0) > 0),
        "rated": len(rated),
        "useful": useful,
        "not_useful": len(rated) - useful,
        # None, пока оценок нет: «0%» ввело бы в заблуждение
        "usefulness": round(useful / len(rated), 3) if rated else None,
    }
    return {
        "period_days": days,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "summary": summary,
        "gaps": [g.to_dict() | {"problem_label": PROBLEM_LABELS[g.problem]} for g in gaps],
        "documents": _document_stats(rows, known_documents or {}),
        "grouping": "embeddings" if used_embeddings else "words",
    }


# ------------------------------ выгрузка в Excel ------------------------------


def build_report_xlsx(report: dict) -> bytes:
    """Отчёт в Excel: сводка, пробелы, документы. Тексты вопросов — недоверенные (см. append_row)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def sheet(title: str, headers: list[str], rows: list[list], widths: list[int], first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        append_row(ws, headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            append_row(ws, row)
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"

    s = report["summary"]
    usefulness = f"{round(s['usefulness'] * 100)}%" if s["usefulness"] is not None else "оценок пока нет"
    sheet(
        "Сводка",
        ["Показатель", "Значение"],
        [
            ["Период, дней", report["period_days"]],
            ["Вопросов к базе знаний", s["questions"]],
            ["Сотрудников задавали вопросы", s["people"]],
            ["Ответов по документам", s["answered"]],
            ["Отказов («в базе нет»)", s["refused"]],
            ["Просьб уточнить вопрос", s["clarifications"]],
            ["Ответов с утверждениями без опоры в документах", s["warned"]],
            ["Оценено пользователями", s["rated"]],
            ["Из них «полезно»", s["useful"]],
            ["Доля «полезно» среди оценённых", usefulness],
        ],
        [52, 22],
        first=True,
    )
    sheet(
        "Пробелы",
        ["Тема (типичный вопрос)", "Проблема", "Сотрудников", "Обращений", "Без ответа", "Не полезно",
         "Подразделения", "Примеры вопросов", "Документы, которые показаны"],
        [
            [
                g["topic"],
                g["problem_label"],
                g["people"],
                g["requests"],
                g["unanswered"],
                g["not_useful"],
                ", ".join(g["divisions"]),
                "\n".join(g["examples"]),
                "\n".join(x["relative_path"] for x in g["sources"]),
            ]
            for g in report["gaps"]
        ],
        [48, 34, 13, 11, 10, 12, 28, 60, 48],
    )
    docs = report["documents"]
    doc_rows = [
        [d["file_name"], "Ответы по документу оценивают как неполезные", d["used"], d["useful"], d["not_useful"]]
        for d in docs["weak"]
    ] + [[d["file_name"], "Не использовался в ответах за период", 0, 0, 0] for d in docs["unused"]]
    sheet(
        "Документы",
        ["Документ", "Замечание", "Использован в ответах", "Полезно", "Не полезно"],
        doc_rows,
        [52, 46, 22, 10, 12],
    )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
